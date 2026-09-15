"""全量 API 接口测试。

覆盖 :mod:`app.routers` 下所有公开接口：
- 公共接口：健康检查、分类、商品列表/详情/搜索
- 认证接口：注册（邮箱验证码）、登录、当前用户
- 购物车接口：查/加/改/删/清空
- 订单接口：下单、订单列表/详情、支付、取消、确认收货
- 后台接口：管理员登录、仪表盘、商品管理、SKU/库存、订单管理、分类、库存流水

运行（需先启动服务）：
    .venv/bin/python -m pytest test/test_api.py -v
"""
from __future__ import annotations

import os
import subprocess
import time
from urllib.parse import urlparse

import pytest
import requests

from app.config import settings

# 目标服务地址：默认 8010，可覆盖（如 export TEST_BASE_URL=http://127.0.0.1:8020）
BASE = os.environ.get("TEST_BASE_URL", "http://127.0.0.1:8010").rstrip("/")
H = {"Accept-Language": "zh"}

BUYER_EMAIL = "buyer@example.com"
BUYER_PASS = "buyer123"
ADMIN_USER = "admin"
ADMIN_PASS = "admin123"

ts = int(time.time())
TEST_EMAIL = f"test_{ts}@example.com"

# 数据库连接参数（psql 子进程读取验证码，避免 async 引擎事件循环冲突）
_DB_URL = urlparse(settings.DATABASE_URL)
_PSQL_ENV = {
    **os.environ,
    "PGHOST": _DB_URL.hostname or "localhost",
    "PGPORT": str(_DB_URL.port or 5432),
    "PGUSER": _DB_URL.username or "",
    "PGDATABASE": (_DB_URL.path or "/").lstrip("/"),
    "PGPASSWORD": _DB_URL.password or "",
}


def get_latest_code(email: str, purpose: str | None = None) -> str:
    """从数据库读取该邮箱最新未使用的验证码（psql 子进程，避免事件循环冲突）"""
    where = f"WHERE email = '{email.lower()}'"
    if purpose:
        where += f" AND purpose = '{purpose}'"
    sql = f"SELECT code FROM email_verify_codes {where} ORDER BY id DESC LIMIT 1;"
    res = subprocess.run(
        ["psql", "-t", "-A", "-c", sql],
        capture_output=True, text=True, env=_PSQL_ENV,
    )
    return res.stdout.strip()


def send_verify_code(email: str, purpose: str = "register") -> requests.Response:
    """请求发送邮箱验证码"""
    return requests.post(
        f"{BASE}/api/auth/send-code",
        json={"email": email, "purpose": purpose},
    )


def register_user(email: str, password: str = "test12345") -> requests.Response:
    """完整注册流程：发送验证码 → 读取 → 注册"""
    send_verify_code(email)
    code = get_latest_code(email)
    return requests.post(
        f"{BASE}/api/auth/register",
        json={"email": email, "password": password, "full_name": "测试用户", "code": code},
    )


# ---------- fixtures ----------
@pytest.fixture(scope="module", autouse=True)
def buyer_setup():
    """确保 buyer 测试账号在首个测试运行前已存在（避免测试顺序依赖）"""
    r = requests.post(f"{BASE}/api/auth/login", json={"email": BUYER_EMAIL, "password": BUYER_PASS})
    if r.status_code != 200:
        rr = register_user(BUYER_EMAIL, BUYER_PASS)
        assert rr.status_code in (200, 201), rr.text


@pytest.fixture(scope="module")
def buyer_token(buyer_setup):
    r = requests.post(f"{BASE}/api/auth/login", json={"email": BUYER_EMAIL, "password": BUYER_PASS})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}", **H}

@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(f"{BASE}/api/admin/login", json={"username": ADMIN_USER, "password": ADMIN_PASS})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}", **H}


# ---------- 公共接口 ----------
class TestPublic:
    def test_health(self):
        r = requests.get(f"{BASE}/api/health")
        assert r.status_code == 200
        assert r.json()["status"] == "ok"

    def test_categories(self):
        r = requests.get(f"{BASE}/api/categories", headers=H)
        assert r.status_code == 200
        cats = r.json()
        assert isinstance(cats, list) and len(cats) > 0
        assert all("id" in c and "name_i18n" in c for c in cats)

    def test_products_list(self):
        r = requests.get(f"{BASE}/api/products", headers=H)
        assert r.status_code == 200
        items = r.json()
        assert isinstance(items, list) and len(items) > 0
        first = items[0]
        for key in ("id", "sku_code", "base_price", "display_name"):
            assert key in first

    def test_products_pagination(self):
        r = requests.get(f"{BASE}/api/products?page=1&page_size=3", headers=H)
        assert r.status_code == 200
        assert len(r.json()) <= 3

    def test_products_featured(self):
        r = requests.get(f"{BASE}/api/products?featured=true", headers=H)
        assert r.status_code == 200
        assert all(p["is_featured"] for p in r.json())

    def test_products_category_filter(self):
        cats = requests.get(f"{BASE}/api/categories", headers=H).json()
        cid = cats[0]["id"]
        r = requests.get(f"{BASE}/api/products", params={"category_id": cid}, headers=H)
        assert r.status_code == 200
        items = r.json()
        assert len(items) > 0
        # 列表响应不含 category_id，改用详情接口核对分类归属
        for p in items[:3]:
            detail = requests.get(f"{BASE}/api/products/{p['id']}", headers=H).json()
            assert detail["category_id"] == cid, f"商品 {p['id']} 不属于分类 {cid}"

    def test_products_search_chinese(self):
        # 中文搜索（JSONB #>> 修复的回归用例）
        r = requests.get(f"{BASE}/api/products", params={"q": "瑜伽"}, headers=H)
        assert r.status_code == 200
        names = [p.get("display_name", "") for p in r.json()]
        assert any("瑜伽" in n or "Yoga" in n for n in names)

    def test_products_search_none(self):
        r = requests.get(f"{BASE}/api/products", params={"q": "不存在的商品xyz"}, headers=H)
        assert r.status_code == 200
        assert r.json() == []

    def test_product_detail(self):
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[0]["id"]
        r = requests.get(f"{BASE}/api/products/{pid}", headers=H)
        assert r.status_code == 200
        detail = r.json()
        assert detail["id"] == pid
        assert "skus" in detail and len(detail["skus"]) >= 1

    def test_product_detail_404(self):
        r = requests.get(f"{BASE}/api/products/999999", headers=H)
        assert r.status_code == 404


# ---------- 订阅 ----------
class TestNewsletter:
    def test_subscribe(self):
        email = f"sub_{ts}@example.com"
        r = requests.post(f"{BASE}/api/subscribe", json={"email": email})
        assert r.status_code == 201, r.text
        assert r.json()["message"]

    def test_subscribe_duplicate(self):
        email = f"sub_{ts}@example.com"
        r = requests.post(f"{BASE}/api/subscribe", json={"email": email})
        assert r.status_code in (200, 201)

    def test_subscribe_invalid_email(self):
        r = requests.post(f"{BASE}/api/subscribe", json={"email": "bad-email"})
        assert r.status_code == 422

    def test_admin_list_subscribers(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/subscribers", headers=admin_token)
        assert r.status_code == 200
        subscribers = r.json()
        assert isinstance(subscribers, list)
        assert any(s["email"] == f"sub_{ts}@example.com" for s in subscribers)

    def test_admin_delete_subscriber(self, admin_token):
        email = f"sub_del_{ts}@example.com"
        requests.post(f"{BASE}/api/subscribe", json={"email": email})
        subs = requests.get(f"{BASE}/api/admin/subscribers", headers=admin_token).json()
        target = next((s for s in subs if s["email"] == email), None)
        if target:
            r = requests.delete(f"{BASE}/api/admin/subscribers/{target['id']}", headers=admin_token)
            assert r.status_code == 200


# ---------- 认证接口 ----------
class TestAuth:
    def test_send_code(self):
        r = send_verify_code(f"code_{ts}@example.com")
        assert r.status_code == 200, r.text
        assert r.json()["message"]

    def test_send_code_existing_email(self):
        r = send_verify_code(BUYER_EMAIL)
        assert r.status_code == 409

    def test_register_requires_code(self):
        payload = {"email": TEST_EMAIL, "password": "test123", "full_name": "测试用户"}
        r = requests.post(f"{BASE}/api/auth/register", json=payload)
        assert r.status_code == 400, r.text

    def test_register_wrong_code(self):
        send_verify_code(TEST_EMAIL)
        payload = {"email": TEST_EMAIL, "password": "test123", "full_name": "测试用户", "code": "000000"}
        r = requests.post(f"{BASE}/api/auth/register", json=payload)
        assert r.status_code == 400, r.text

    def test_register(self):
        send_verify_code(TEST_EMAIL)
        code = get_latest_code(TEST_EMAIL)
        assert code, "验证码未写入数据库"
        payload = {"email": TEST_EMAIL, "password": "test123", "full_name": "测试用户", "code": code}
        r = requests.post(f"{BASE}/api/auth/register", json=payload)
        assert r.status_code == 201, r.text
        assert "access_token" in r.json()

    def test_register_duplicate(self):
        payload = {"email": BUYER_EMAIL, "password": "whatever123", "code": "000000"}
        r = requests.post(f"{BASE}/api/auth/register", json=payload)
        assert r.status_code == 409

    def test_login_ok(self):
        r = requests.post(f"{BASE}/api/auth/login", json={"email": BUYER_EMAIL, "password": BUYER_PASS})
        assert r.status_code == 200
        assert "access_token" in r.json()

    def test_login_wrong_password(self):
        r = requests.post(f"{BASE}/api/auth/login", json={"email": BUYER_EMAIL, "password": "wrong"})
        assert r.status_code == 401

    def test_login_missing_user(self):
        r = requests.post(f"{BASE}/api/auth/login", json={"email": "nobody@example.com", "password": "x123456"})
        assert r.status_code == 401

    def test_me(self, buyer_token):
        r = requests.get(f"{BASE}/api/auth/me", headers=buyer_token)
        assert r.status_code == 200
        assert r.json()["email"] == BUYER_EMAIL

    def test_me_unauthorized(self):
        r = requests.get(f"{BASE}/api/auth/me")
        assert r.status_code == 401

    # ---------- 重置密码 ----------
    def test_reset_send_code_unknown_email(self):
        r = send_verify_code(f"nobody_{ts}@example.com", purpose="reset")
        assert r.status_code == 404

    def test_reset_password_wrong_code(self):
        send_verify_code(TEST_EMAIL, purpose="reset")
        r = requests.post(f"{BASE}/api/auth/reset-password",
                          json={"email": TEST_EMAIL, "code": "000000", "new_password": "newpass123"})
        assert r.status_code == 400

    def test_reset_password_flow(self):
        # 使用 TEST_EMAIL 注册的账号（test_register 已创建）
        send_verify_code(TEST_EMAIL, purpose="reset")
        code = get_latest_code(TEST_EMAIL, purpose="reset")
        assert code, "reset 验证码未写入"
        r = requests.post(f"{BASE}/api/auth/reset-password",
                          json={"email": TEST_EMAIL, "code": code, "new_password": "resetpass123"})
        assert r.status_code == 200, r.text
        # 新密码可登录
        login_ok = requests.post(f"{BASE}/api/auth/login",
                                 json={"email": TEST_EMAIL, "password": "resetpass123"})
        assert login_ok.status_code == 200
        # 旧密码失效
        login_old = requests.post(f"{BASE}/api/auth/login",
                                  json={"email": TEST_EMAIL, "password": "test123"})
        assert login_old.status_code == 401

    def test_reset_password_reuse_code(self):
        # 验证码只能使用一次
        send_verify_code(TEST_EMAIL, purpose="reset")
        code = get_latest_code(TEST_EMAIL, purpose="reset")
        r1 = requests.post(f"{BASE}/api/auth/reset-password",
                           json={"email": TEST_EMAIL, "code": code, "new_password": "againpass123"})
        assert r1.status_code == 200
        r2 = requests.post(f"{BASE}/api/auth/reset-password",
                           json={"email": TEST_EMAIL, "code": code, "new_password": "hackpass123"})
        assert r2.status_code == 400


# ---------- 购物车接口 ----------
class TestCart:
    def _first_sku(self):
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[0]["id"]
        detail = requests.get(f"{BASE}/api/products/{pid}", headers=H).json()
        return detail["skus"][0]["id"], detail["skus"][0]["available_stock"]

    def test_cart_requires_login(self):
        r = requests.get(f"{BASE}/api/cart")
        assert r.status_code == 401

    def test_add_and_get(self, buyer_token):
        sku_id, _ = self._first_sku()
        r = requests.post(f"{BASE}/api/cart/items", json={"sku_id": sku_id, "quantity": 1}, headers=buyer_token)
        assert r.status_code == 201, r.text
        cart = r.json()
        assert len(cart["items"]) >= 1
        assert str(cart["total_amount"]) != "0"

        # 获取购物车
        r2 = requests.get(f"{BASE}/api/cart", headers=buyer_token)
        assert r2.status_code == 200

    def test_add_invalid_quantity(self, buyer_token):
        sku_id, _ = self._first_sku()
        r = requests.post(f"{BASE}/api/cart/items", json={"sku_id": sku_id, "quantity": 0}, headers=buyer_token)
        assert r.status_code == 422

    def test_add_nonexistent_sku(self, buyer_token):
        r = requests.post(f"{BASE}/api/cart/items", json={"sku_id": 999999, "quantity": 1}, headers=buyer_token)
        assert r.status_code == 404

    def test_update_and_remove(self, buyer_token):
        # 加购
        sku_id, stock = self._first_sku()
        add = requests.post(f"{BASE}/api/cart/items", json={"sku_id": sku_id, "quantity": 2}, headers=buyer_token)
        assert add.status_code == 201
        item_id = add.json()["items"][0]["id"]

        # 修改数量
        upd = requests.put(f"{BASE}/api/cart/items/{item_id}", json={"quantity": 3}, headers=buyer_token)
        assert upd.status_code == 200
        assert upd.json()["items"][0]["quantity"] == 3

        # 库存上限校验
        over_qty = stock + 5 if stock < 1_000_000 else 500
        r_over = requests.put(f"{BASE}/api/cart/items/{item_id}", json={"quantity": over_qty}, headers=buyer_token)
        assert r_over.status_code == 400

        # 删除
        rem = requests.delete(f"{BASE}/api/cart/items/{item_id}", headers=buyer_token)
        assert rem.status_code == 200

        # 删除不存在
        r_404 = requests.delete(f"{BASE}/api/cart/items/{item_id}", headers=buyer_token)
        assert r_404.status_code == 404

    def test_clear(self, buyer_token):
        sku_id, _ = self._first_sku()
        requests.post(f"{BASE}/api/cart/items", json={"sku_id": sku_id, "quantity": 1}, headers=buyer_token)
        r = requests.delete(f"{BASE}/api/cart", headers=buyer_token)
        assert r.status_code == 200
        cart = requests.get(f"{BASE}/api/cart", headers=buyer_token).json()
        assert cart["items"] == []


# ---------- 高级筛选 / Autocomplete ----------
class TestCatalogFilter:
    def test_min_max_price(self):
        r = requests.get(f"{BASE}/api/products", params={"min_price": 100, "max_price": 5000}, headers=H)
        assert r.status_code == 200
        items = r.json()
        assert len(items) > 0
        for p in items:
            assert 100 <= float(p["base_price"]) <= 5000

    def test_min_price_only(self):
        r = requests.get(f"{BASE}/api/products", params={"min_price": 99999}, headers=H)
        assert r.status_code == 200
        assert r.json() == []

    def test_sort_price_desc(self):
        r = requests.get(f"{BASE}/api/products", params={"sort_by": "price_desc"}, headers=H)
        assert r.status_code == 200
        prices = [float(p["base_price"]) for p in r.json()]
        assert prices == sorted(prices, reverse=True)

    def test_sort_sales_desc(self):
        r = requests.get(f"{BASE}/api/products", params={"sort_by": "sales_desc"}, headers=H)
        assert r.status_code == 200
        sales = [p.get("sales_count", 0) for p in r.json()]
        assert sales == sorted(sales, reverse=True)

    def test_autocomplete_zh(self):
        r = requests.get(f"{BASE}/api/products/autocomplete", params={"q": "瑜伽"}, headers=H)
        assert r.status_code == 200
        results = r.json()
        assert isinstance(results, list) and len(results) > 0
        assert all("id" in x and "name_zh" in x for x in results)

    def test_autocomplete_en(self):
        # 真实种子商品英文名含 bra/tank/top，不含 yoga
        r = requests.get(f"{BASE}/api/products/autocomplete", params={"q": "bra"}, headers=H)
        assert r.status_code == 200
        assert len(r.json()) > 0

    def test_autocomplete_empty(self):
        r = requests.get(f"{BASE}/api/products/autocomplete", params={"q": "zzz不存在"}, headers=H)
        assert r.status_code == 200
        assert r.json() == []


# ---------- 收藏 ----------
class TestWishlist:
    def test_requires_auth(self):
        r = requests.get(f"{BASE}/api/wishlist")
        assert r.status_code == 401

    def test_add_and_list(self, buyer_token):
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[0]["id"]
        r = requests.post(f"{BASE}/api/wishlist", json={"product_id": pid}, headers=buyer_token)
        assert r.status_code == 201, r.text
        data = r.json()
        assert data["total"] >= 1
        assert pid in data["product_ids"]

        lst = requests.get(f"{BASE}/api/wishlist", headers=buyer_token).json()
        assert any(i["product_id"] == pid for i in lst["items"])

    def test_add_duplicate_is_idempotent(self, buyer_token):
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[0]["id"]
        requests.post(f"{BASE}/api/wishlist", json={"product_id": pid}, headers=buyer_token)
        r2 = requests.post(f"{BASE}/api/wishlist", json={"product_id": pid}, headers=buyer_token)
        # 幂等：再次添加不报错（201 或 200 均可），且列表中只有一条
        assert r2.status_code in (200, 201)
        lst = requests.get(f"{BASE}/api/wishlist", headers=buyer_token).json()
        count = sum(1 for i in lst["items"] if i["product_id"] == pid)
        assert count == 1

    def test_remove(self, buyer_token):
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[1 % len(items)]["id"]
        requests.post(f"{BASE}/api/wishlist", json={"product_id": pid}, headers=buyer_token)
        r = requests.delete(f"{BASE}/api/wishlist/{pid}", headers=buyer_token)
        assert r.status_code == 200
        lst = requests.get(f"{BASE}/api/wishlist", headers=buyer_token).json()
        assert all(i["product_id"] != pid for i in lst["items"])

    def test_move_to_cart(self, buyer_token):
        # 清空购物车再测试
        requests.delete(f"{BASE}/api/cart", headers=buyer_token)
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[0]["id"]
        detail = requests.get(f"{BASE}/api/products/{pid}", headers=H).json()
        sku_id = detail["skus"][0]["id"]
        requests.post(f"{BASE}/api/wishlist", json={"product_id": pid}, headers=buyer_token)
        r = requests.post(f"{BASE}/api/wishlist/move-to-cart/{pid}", headers=buyer_token)
        assert r.status_code == 200, r.text
        cart = requests.get(f"{BASE}/api/cart", headers=buyer_token).json()
        assert any(c["sku_id"] == sku_id for c in cart["items"])
        # 移入购物车后应取消收藏
        lst = requests.get(f"{BASE}/api/wishlist", headers=buyer_token).json()
        assert all(i["product_id"] != pid for i in lst["items"])

    def test_clear(self, buyer_token):
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        for p in items[:2]:
            requests.post(f"{BASE}/api/wishlist", json={"product_id": p["id"]}, headers=buyer_token)
        r = requests.delete(f"{BASE}/api/wishlist", headers=buyer_token)
        assert r.status_code == 200
        lst = requests.get(f"{BASE}/api/wishlist", headers=buyer_token).json()
        assert lst["items"] == []

    def test_admin_list_exposes_favorite_count(self, buyer_token, admin_token):
        """后台商品列表必须返回 favorite_count（否则「收藏」列恒为 0）

        `favorite_count` 不是 products 表字段，需要在路由里聚合 WishlistItem。
        """
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[0]["id"]
        requests.delete(f"{BASE}/api/wishlist", headers=buyer_token)
        requests.post(f"{BASE}/api/wishlist", json={"product_id": pid}, headers=buyer_token)

        got = requests.get(f"{BASE}/api/admin/products?product_id={pid}",
                           headers=admin_token).json()
        row = next(p for p in got if p["id"] == pid)
        assert "favorite_count" in row
        assert row["favorite_count"] == 1

        # 取消收藏后归零
        requests.delete(f"{BASE}/api/wishlist/{pid}", headers=buyer_token)
        got2 = requests.get(f"{BASE}/api/admin/products?product_id={pid}",
                            headers=admin_token).json()
        row2 = next(p for p in got2 if p["id"] == pid)
        assert row2["favorite_count"] == 0


# ---------- 地址 ----------
class TestAddresses:
    def test_requires_auth(self):
        r = requests.get(f"{BASE}/api/addresses")
        assert r.status_code == 401

    def test_crud(self, buyer_token):
        # 创建
        r = requests.post(f"{BASE}/api/addresses", json={
            "receiver_name": "测试收货人", "receiver_phone": "13900000000",
            "detail": "上海市浦东新区测试路 88 号", "is_default": True,
        }, headers=buyer_token)
        assert r.status_code == 201, r.text
        addr = r.json()
        assert addr["is_default"] is True
        aid = addr["id"]

        # 列表
        lst = requests.get(f"{BASE}/api/addresses", headers=buyer_token).json()
        assert any(a["id"] == aid for a in lst)

        # 更新
        r = requests.put(f"{BASE}/api/addresses/{aid}", json={"receiver_name": "新名字"}, headers=buyer_token)
        assert r.status_code == 200
        assert r.json()["receiver_name"] == "新名字"

        # 删除
        r = requests.delete(f"{BASE}/api/addresses/{aid}", headers=buyer_token)
        assert r.status_code == 200
        lst2 = requests.get(f"{BASE}/api/addresses", headers=buyer_token).json()
        assert all(a["id"] != aid for a in lst2)

    def test_default_switch(self, buyer_token):
        a1 = requests.post(f"{BASE}/api/addresses", json={
            "receiver_name": "甲", "receiver_phone": "139",
            "detail": "地址一", "is_default": False,
        }, headers=buyer_token).json()
        a2 = requests.post(f"{BASE}/api/addresses", json={
            "receiver_name": "乙", "receiver_phone": "138",
            "detail": "地址二", "is_default": True,
        }, headers=buyer_token).json()
        # 新默认生效，之前的取消默认
        lst = requests.get(f"{BASE}/api/addresses", headers=buyer_token).json()
        for a in lst:
            if a["id"] == a2["id"]:
                assert a["is_default"] is True
            if a["id"] == a1["id"]:
                assert a["is_default"] is False
        # 清理
        requests.delete(f"{BASE}/api/addresses/{a1['id']}", headers=buyer_token)
        requests.delete(f"{BASE}/api/addresses/{a2['id']}", headers=buyer_token)

    def test_default_first_order(self, buyer_token):
        """默认地址排在最前"""
        requests.post(f"{BASE}/api/addresses", json={
            "receiver_name": "普通", "receiver_phone": "137", "detail": "普通地址",
        }, headers=buyer_token)
        addr2 = requests.post(f"{BASE}/api/addresses", json={
            "receiver_name": "默认主", "receiver_phone": "136", "detail": "默认地址", "is_default": True,
        }, headers=buyer_token).json()
        lst = requests.get(f"{BASE}/api/addresses", headers=buyer_token).json()
        assert lst[0]["id"] == addr2["id"]
        for a in lst:
            requests.delete(f"{BASE}/api/addresses/{a['id']}", headers=buyer_token)


# ---------- 评价 ----------
class TestReviews:
    def _new_user_token(self):
        """注册一个全新的临时用户（避免历史购买/评价数据干扰）"""
        email = f"rv_{ts}_{int(time.time()*1000)}@example.com"
        r = register_user(email)
        assert r.status_code == 201, r.text
        return {"Authorization": f"Bearer {r.json()['access_token']}", **H}

    def test_list_empty_public(self):
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[0]["id"]
        r = requests.get(f"{BASE}/api/products/{pid}/reviews", headers=H)
        assert r.status_code == 200
        assert isinstance(r.json(), list)

    def test_review_requires_purchase(self, admin_token):
        """未购买过的商品不能评价（新用户无任何购买记录）"""
        token = self._new_user_token()
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[-1]["id"]
        r = requests.post(f"{BASE}/api/products/{pid}/reviews", json={
            "rating": 5, "title": "不错", "content": "很喜欢这个商品",
        }, headers=token)
        assert r.status_code == 403, r.text

    def test_review_requires_auth(self):
        r = requests.post(f"{BASE}/api/products/1/reviews", json={
            "rating": 5, "title": "t", "content": "c",
        })
        assert r.status_code == 401

    def test_review_not_approved_until_moderation(self, admin_token):
        """评价需审核后才能公开可见"""
        # 新用户：注册 -> 下单 -> 评价 -> 审核前不可见 -> 审核后可见
        token = self._new_user_token()
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        pid = items[-1]["id"]

        # 1. 未购买时不可评
        can = requests.get(f"{BASE}/api/products/{pid}/reviewable", headers=token).json()
        assert can.get("reviewable") is False

        # 2. 下单购买
        detail = requests.get(f"{BASE}/api/products/{pid}", headers=H).json()
        sku_id = detail["skus"][0]["id"]
        requests.post(f"{BASE}/api/cart/items", json={"sku_id": sku_id, "quantity": 1}, headers=token)
        r = requests.post(f"{BASE}/api/orders/checkout", json={
            "receiver_name": "评价测试", "receiver_phone": "138",
            "receiver_address": "测试地址", "payment_method": "mock",
        }, headers=token)
        assert r.status_code == 201, r.text

        # 3. 购买后可评
        can2 = requests.get(f"{BASE}/api/products/{pid}/reviewable", headers=token).json()
        assert can2.get("reviewable") is True

        # 4. 提交评价 -> pending
        r = requests.post(f"{BASE}/api/products/{pid}/reviews", json={
            "rating": 4, "title": "审核测试", "content": "这条评价待审核",
        }, headers=token)
        assert r.status_code == 201, r.text
        review_id = r.json()["id"]

        # 5. 审核前公开列表不可见
        lst = requests.get(f"{BASE}/api/products/{pid}/reviews", headers=H).json()
        assert all(x["id"] != review_id for x in lst)

        # 6. 重复评价被拒（每人一评）
        r = requests.post(f"{BASE}/api/products/{pid}/reviews", json={
            "rating": 4, "title": "重复", "content": "重复评价",
        }, headers=token)
        assert r.status_code == 409, r.text

        # 7. 管理员审核通过
        r = requests.post(f"{BASE}/api/admin/reviews/{review_id}/moderate",
                          json={"status": "approved"}, headers=admin_token)
        assert r.status_code == 200, r.text

        # 8. 审核后公开可见
        lst2 = requests.get(f"{BASE}/api/products/{pid}/reviews", headers=H).json()
        assert any(x["id"] == review_id for x in lst2)

        # 9. 我的评价列表可见
        mine = requests.get(f"{BASE}/api/my-reviews", headers=token).json()
        assert any(x["id"] == review_id for x in mine)

    def test_reviewable_requires_auth(self):
        r = requests.get(f"{BASE}/api/products/1/reviewable")
        assert r.status_code == 401


# ---------- 用户故事 ----------
class TestUserStories:
    def test_upload_edit_moderate_and_delete(self, buyer_token, admin_token):
        """用户故事完整流程：图片上传 -> 投稿 -> 本人编辑 -> 管理员发布 -> 公开 -> 删除。"""
        image = requests.post(
            f"{BASE}/api/user-story-upload",
            files={"file": ("story.png", b"fake-png-content", "image/png")},
            headers=buyer_token,
        )
        assert image.status_code == 201, image.text
        image_url = image.json()["url"]

        created = requests.post(f"{BASE}/api/user-stories", json={
            "title": "我的初次分享",
            "content": "这是我的用户故事。",
            "image_url": image_url,
        }, headers=buyer_token)
        assert created.status_code == 201, created.text
        story_id = created.json()["id"]
        assert created.json()["is_published"] is False
        public_before = requests.get(f"{BASE}/api/user-stories").json()
        public_items = public_before["items"] if isinstance(public_before, dict) else public_before
        assert not any(s["id"] == story_id for s in public_items)

        mine = requests.get(f"{BASE}/api/user-stories/mine", headers=buyer_token)
        assert mine.status_code == 200
        assert mine.json()[0]["image_url"] == image_url

        edited = requests.put(f"{BASE}/api/user-stories/{story_id}", json={
            "title": "编辑后的分享",
            "content": "我补充了更多故事内容。",
            "image_url": image_url,
        }, headers=buyer_token)
        assert edited.status_code == 200, edited.text
        assert edited.json()["title"] == "编辑后的分享"
        assert edited.json()["is_published"] is False

        published = requests.put(f"{BASE}/api/admin/user-stories/{story_id}", json={
            "is_published": True,
        }, headers=admin_token)
        assert published.status_code == 200, published.text
        public_resp = requests.get(f"{BASE}/api/user-stories").json()
        public = public_resp["items"] if isinstance(public_resp, dict) else public_resp
        public_story = next(s for s in public if s["id"] == story_id)
        assert public_story["content"] == "我补充了更多故事内容。"
        assert public_story["image_url"] == image_url

        removed = requests.delete(f"{BASE}/api/admin/user-stories/{story_id}", headers=admin_token)
        assert removed.status_code == 200, removed.text

    def test_user_story_requires_auth(self):
        assert requests.post(f"{BASE}/api/user-stories", json={"title": "t", "content": "c"}).status_code == 401
        assert requests.get(f"{BASE}/api/user-stories/mine").status_code == 401

    def test_public_detail_endpoint(self, buyer_token, admin_token):
        """公开详情接口：仅已发布可见；未发布/不存在返回 404。"""
        created = requests.post(f"{BASE}/api/user-stories", json={
            "title": "详情接口测试",
            "content": "详情内容。",
        }, headers=buyer_token)
        assert created.status_code == 201, created.text
        story_id = created.json()["id"]

        # 未发布时详情不可见
        detail_pending = requests.get(f"{BASE}/api/user-stories/{story_id}")
        assert detail_pending.status_code == 404

        # 管理员发布后可见
        requests.put(f"{BASE}/api/admin/user-stories/{story_id}", json={"is_published": True}, headers=admin_token)
        detail = requests.get(f"{BASE}/api/user-stories/{story_id}")
        assert detail.status_code == 200, detail.text
        assert detail.json()["id"] == story_id
        assert detail.json()["title"] == "详情接口测试"

        # 不存在的 id
        assert requests.get(f"{BASE}/api/user-stories/999999999").status_code == 404
        # 清理
        requests.delete(f"{BASE}/api/admin/user-stories/{story_id}", headers=admin_token)

    def test_user_delete_own_story(self, buyer_token, admin_token):
        """用户可删除自己的故事；删除后公开列表和详情均不可见。"""
        created = requests.post(f"{BASE}/api/user-stories", json={
            "title": "待删除的故事",
            "content": "这段内容将被删除。",
        }, headers=buyer_token)
        assert created.status_code == 201, created.text
        story_id = created.json()["id"]

        # 非本人不可删（换一个临时用户 token 无法轻易获得，仅验证无 token 401 与本人删除）
        assert requests.delete(f"{BASE}/api/user-stories/{story_id}").status_code == 401

        deleted = requests.delete(f"{BASE}/api/user-stories/{story_id}", headers=buyer_token)
        assert deleted.status_code == 200, deleted.text

        mine = requests.get(f"{BASE}/api/user-stories/mine", headers=buyer_token).json()
        assert not any(s["id"] == story_id for s in mine)
        assert requests.get(f"{BASE}/api/user-stories/{story_id}").status_code == 404

    def test_like_and_comment_flow(self, buyer_token, admin_token):
        """点赞/评论完整流程：发表故事 -> 发布 -> 点赞幂等 -> 评论 -> 计数更新 -> 取消点赞。"""
        created = requests.post(f"{BASE}/api/user-stories", json={
            "title": "点赞评论测试故事",
            "content": "等待被点赞和评论。",
        }, headers=buyer_token)
        assert created.status_code == 201, created.text
        story_id = created.json()["id"]
        requests.put(f"{BASE}/api/admin/user-stories/{story_id}", json={"is_published": True}, headers=admin_token)

        # 未登录不能点赞/评论
        assert requests.post(f"{BASE}/api/user-stories/{story_id}/like").status_code == 401
        assert requests.post(f"{BASE}/api/user-stories/{story_id}/comments", json={"content": "x"}).status_code == 401

        # 点赞（幂等）
        like1 = requests.post(f"{BASE}/api/user-stories/{story_id}/like", headers=buyer_token)
        assert like1.status_code == 200, like1.text
        assert like1.json() == {"liked": True, "like_count": 1}
        like2 = requests.post(f"{BASE}/api/user-stories/{story_id}/like", headers=buyer_token)
        assert like2.json() == {"liked": True, "like_count": 1}

        # 评论
        cm = requests.post(f"{BASE}/api/user-stories/{story_id}/comments", json={"content": "很棒！"}, headers=buyer_token)
        assert cm.status_code == 200, cm.text
        assert cm.json()["content"] == "很棒！"
        # 空评论 400
        assert requests.post(f"{BASE}/api/user-stories/{story_id}/comments", json={"content": "  "}, headers=buyer_token).status_code == 400

        # 评论列表（author 是注册时填的 full_name）
        cl = requests.get(f"{BASE}/api/user-stories/{story_id}/comments").json()
        assert len(cl) == 1 and cl[0]["author"] == "测试用户"

        # 详情带计数与 liked 状态（带 token 时 liked 应为 True）
        detail = requests.get(f"{BASE}/api/user-stories/{story_id}").json()
        assert detail["like_count"] == 1
        assert detail["comment_count"] == 1
        # 未登录：liked 恒为 False
        assert detail["liked"] is False
        detail_auth = requests.get(f"{BASE}/api/user-stories/{story_id}", headers=buyer_token).json()
        assert detail_auth["liked"] is True

        # 公开列表带计数
        pub = requests.get(f"{BASE}/api/user-stories").json()
        items = pub["items"] if isinstance(pub, dict) else pub
        pub_story = next(s for s in items if s["id"] == story_id)
        assert pub_story["like_count"] == 1
        assert pub_story["comment_count"] == 1

        # 取消点赞
        ul = requests.post(f"{BASE}/api/user-stories/{story_id}/unlike", headers=buyer_token)
        assert ul.json() == {"liked": False, "like_count": 0}
        assert requests.get(f"{BASE}/api/user-stories/{story_id}").json()["liked"] is False

        # 清理
        requests.delete(f"{BASE}/api/admin/user-stories/{story_id}", headers=admin_token)


# ---------- 订单与支付接口 ----------
class TestOrders:
    @pytest.fixture()
    def order_no(self, buyer_token):
        # 准备：清空购物车，加购一个商品，下单
        requests.delete(f"{BASE}/api/cart", headers=buyer_token)
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        detail = requests.get(f"{BASE}/api/products/{items[0]['id']}", headers=H).json()
        sku_id = detail["skus"][0]["id"]
        requests.post(f"{BASE}/api/cart/items", json={"sku_id": sku_id, "quantity": 1}, headers=buyer_token)

        payload = {
            "receiver_name": "张三",
            "receiver_phone": "13800138000",
            "receiver_address": "北京市朝阳区测试路 1 号",
            "remark": "pytest 下单",
            "payment_method": "mock",
        }
        r = requests.post(f"{BASE}/api/orders/checkout", json=payload, headers=buyer_token)
        assert r.status_code == 201, r.text
        return r.json()["order_no"]

    def test_checkout_empty_cart(self, buyer_token):
        requests.delete(f"{BASE}/api/cart", headers=buyer_token)
        payload = {
            "receiver_name": "张三", "receiver_phone": "138", "receiver_address": "测试",
            "payment_method": "mock",
        }
        r = requests.post(f"{BASE}/api/orders/checkout", json=payload, headers=buyer_token)
        assert r.status_code == 400

    def test_checkout_bad_method(self, buyer_token):
        requests.delete(f"{BASE}/api/cart", headers=buyer_token)
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        detail = requests.get(f"{BASE}/api/products/{items[0]['id']}", headers=H).json()
        requests.post(f"{BASE}/api/cart/items",
                      json={"sku_id": detail["skus"][0]["id"], "quantity": 1}, headers=buyer_token)
        payload = {
            "receiver_name": "张三", "receiver_phone": "138", "receiver_address": "测试",
            "payment_method": "alipay",  # 未开通通道
        }
        r = requests.post(f"{BASE}/api/orders/checkout", json=payload, headers=buyer_token)
        assert r.status_code == 400
        assert "尚未开通" in r.json()["detail"] or "未配置" in r.json()["detail"]

    def test_checkout_missing_fields(self, buyer_token):
        r = requests.post(f"{BASE}/api/orders/checkout", json={"payment_method": "mock"}, headers=buyer_token)
        assert r.status_code in (400, 422)

    def test_order_list_and_detail(self, order_no, buyer_token):
        r = requests.get(f"{BASE}/api/orders", headers=buyer_token)
        assert r.status_code == 200
        orders = r.json()
        assert any(o["order_no"] == order_no for o in orders)

        r2 = requests.get(f"{BASE}/api/orders/{order_no}", headers=buyer_token)
        assert r2.status_code == 200
        assert r2.json()["order_no"] == order_no

    def test_checkout_reduces_available_stock_and_cancel_restores(self, buyer_token):
        """下单预占可售库存，取消订单后释放预占库存。"""
        requests.delete(f"{BASE}/api/cart", headers=buyer_token)
        products = requests.get(f"{BASE}/api/products", headers=H).json()
        detail = requests.get(f"{BASE}/api/products/{products[0]['id']}", headers=H).json()
        sku_id = detail["skus"][0]["id"]
        before = detail["skus"][0]["available_stock"]
        assert before >= 1

        add = requests.post(
            f"{BASE}/api/cart/items",
            json={"sku_id": sku_id, "quantity": 1},
            headers=buyer_token,
        )
        assert add.status_code == 201, add.text
        checkout = requests.post(
            f"{BASE}/api/orders/checkout",
            json={
                "receiver_name": "库存测试",
                "receiver_phone": "13800000000",
                "receiver_address": "测试地址",
                "payment_method": "mock",
            },
            headers=buyer_token,
        )
        assert checkout.status_code == 201, checkout.text

        after_checkout = requests.get(f"{BASE}/api/products/{products[0]['id']}", headers=H).json()
        assert after_checkout["skus"][0]["available_stock"] == before - 1

        order_no = checkout.json()["order_no"]
        cancel = requests.post(f"{BASE}/api/orders/{order_no}/cancel", headers=buyer_token)
        assert cancel.status_code == 200, cancel.text

        after_cancel = requests.get(f"{BASE}/api/products/{products[0]['id']}", headers=H).json()
        assert after_cancel["skus"][0]["available_stock"] == before

    def test_order_flow_pay_confirm(self, order_no, buyer_token, admin_token):
        # 支付
        r = requests.post(f"{BASE}/api/orders/{order_no}/pay", headers=buyer_token)
        assert r.status_code == 200, r.text
        pay = r.json()
        assert pay["success"] is True
        assert "transaction_no" in pay and "pay_url" in pay

        # 模拟支付确认
        txn = pay["transaction_no"]
        r2 = requests.get(f"{BASE}/api/payments/mock/confirm", params={"txn_no": txn})
        assert r2.status_code == 200, r2.text
        assert r2.json()["success"] is True
        assert r2.json()["order_no"] == order_no

        # 订单应为已支付
        od = requests.get(f"{BASE}/api/orders/{order_no}", headers=buyer_token).json()
        assert od["status"] == "paid"

        # 管理后台发货
        ship = requests.post(f"{BASE}/api/admin/orders/{order_no}/ship", headers=admin_token)
        assert ship.status_code == 200, ship.text

        # 客户确认收货
        conf = requests.post(f"{BASE}/api/orders/{order_no}/confirm", headers=buyer_token)
        assert conf.status_code == 200, conf.text

        od2 = requests.get(f"{BASE}/api/orders/{order_no}", headers=buyer_token).json()
        assert od2["status"] == "completed"

    def test_order_cancel(self, buyer_token):
        # 准备订单
        requests.delete(f"{BASE}/api/cart", headers=buyer_token)
        items = requests.get(f"{BASE}/api/products", headers=H).json()
        detail = requests.get(f"{BASE}/api/products/{items[0]['id']}", headers=H).json()
        requests.post(f"{BASE}/api/cart/items",
                      json={"sku_id": detail["skus"][0]["id"], "quantity": 1}, headers=buyer_token)
        payload = {
            "receiver_name": "李四", "receiver_phone": "139", "receiver_address": "测试",
            "payment_method": "mock",
        }
        r = requests.post(f"{BASE}/api/orders/checkout", json=payload, headers=buyer_token)
        assert r.status_code == 201
        order_no = r.json()["order_no"]

        r2 = requests.post(f"{BASE}/api/orders/{order_no}/cancel", headers=buyer_token)
        assert r2.status_code == 200, r2.text
        od = requests.get(f"{BASE}/api/orders/{order_no}", headers=buyer_token).json()
        assert od["status"] == "cancelled"

    def test_pay_unknown_order(self, buyer_token):
        r = requests.post(f"{BASE}/api/orders/ORD_NOT_EXIST/pay", headers=buyer_token)
        assert r.status_code == 404

    def test_cancel_unknown_order(self, buyer_token):
        r = requests.post(f"{BASE}/api/orders/ORD_NOT_EXIST/cancel", headers=buyer_token)
        assert r.status_code == 404

    def test_payment_callback_nonexistent(self):
        # mock 网关信任任意 txn_no，未知交易号由 _mark_payment_success 抛 404
        r = requests.post(f"{BASE}/api/payments/callback", json={"txn_no": "NO_SUCH_TXN"})
        assert r.status_code == 404

    def test_payment_query_nonexistent(self):
        # 真正的 404 场景在查询接口
        r = requests.get(f"{BASE}/api/payments/query/NO_SUCH_TXN")
        assert r.status_code == 404


# ---------- 后台接口 ----------
class TestAdmin:
    def test_admin_login_bad(self):
        r = requests.post(f"{BASE}/api/admin/login", json={"username": ADMIN_USER, "password": "wrong"})
        assert r.status_code == 401

    def test_dashboard(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/dashboard", headers=admin_token)
        assert r.status_code == 200, r.text
        data = r.json()
        for key in ("products_count", "orders_count", "customers_count", "revenue"):
            assert key in data

    def test_admin_list_products(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/products", headers=admin_token)
        assert r.status_code == 200
        assert len(r.json()) >= 1

    def test_admin_list_orders(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/orders", headers=admin_token)
        assert r.status_code == 200

    def test_admin_list_orders_status_filter(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/orders", params={"status": "pending"}, headers=admin_token)
        assert r.status_code == 200

    def test_admin_categories(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/categories", headers=admin_token)
        assert r.status_code == 200
        names = {c["name_i18n"].get("zh") for c in r.json()}
        # 真实种子分类（无 demo 的瑜伽套装/上衣）
        assert {"文胸", "背心", "短袖", "长袖", "外套", "短裤", "长裤", "裙子", "其他"}.issubset(names)

    def test_admin_stock_movements(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/stock-movements", headers=admin_token)
        assert r.status_code == 200

    def test_admin_requires_auth(self):
        r = requests.get(f"{BASE}/api/admin/dashboard")
        assert r.status_code == 401

    def test_customer_token_rejected_on_admin(self, buyer_token):
        # 客户 token 不能访问后台
        r = requests.get(f"{BASE}/api/admin/dashboard", headers=buyer_token)
        assert r.status_code in (401, 403)

    def test_admin_create_product_and_sku(self, admin_token):
        sku_code = f"PYTEST-{ts}"
        payload = {
            "sku_code": sku_code,
            "name_zh": f"测试商品 {ts}",
            "name_en": f"Test Product {ts}",
            "description_zh": "pytest 创建",
            "base_price": "123.45",
            "status": "active",
            "skus": [
                {"sku_code": f"{sku_code}-S", "attributes": {"颜色": "红色"}, "price": "123.45", "stock": 10},
            ],
        }
        r = requests.post(f"{BASE}/api/admin/products", json=payload, headers=admin_token)
        assert r.status_code == 201, r.text
        product = r.json()
        assert product["sku_code"] == sku_code
        assert len(product["skus"]) == 1

        # 为商品新增 SKU
        sku_in = {"sku_code": f"{sku_code}-XL2", "attributes": {"颜色": "蓝色"},
                  "price": "130.00", "stock": 5, "is_active": True}
        r2 = requests.post(f"{BASE}/api/admin/products/{product['id']}/skus", json=sku_in, headers=admin_token)
        assert r2.status_code == 201, r2.text

        # 库存调整
        r3 = requests.put(f"{BASE}/api/admin/skus/{product['skus'][0]['id']}/stock",
                          json={"stock": 20, "reason": "pytest_adjust"}, headers=admin_token)
        assert r3.status_code == 200, r3.text
        assert r3.json()["stock"] == 20

    def test_admin_create_category(self, admin_token):
        code = f"pytest-{ts}"
        payload = {"code": code, "name_i18n": {"zh": f"测试分类 {ts}", "en": f"Test Cat {ts}"},
                   "sort_order": 999, "is_active": True}
        r = requests.post(f"{BASE}/api/admin/categories", json=payload, headers=admin_token)
        assert r.status_code == 201, r.text
        category = r.json()
        assert category["code"] == code

        updated = requests.put(f"{BASE}/api/admin/categories/{category['id']}", json={
            "name_i18n": {"zh": "更新分类", "en": "Updated Category"},
        }, headers=admin_token)
        assert updated.status_code == 200, updated.text
        assert updated.json()["name_i18n"]["zh"] == "更新分类"

        duplicate = requests.post(f"{BASE}/api/admin/categories", json=payload, headers=admin_token)
        assert duplicate.status_code == 409

        deleted = requests.delete(f"{BASE}/api/admin/categories/{category['id']}", headers=admin_token)
        assert deleted.status_code == 200, deleted.text

    def test_admin_cannot_delete_category_in_use(self, admin_token):
        products = requests.get(f"{BASE}/api/admin/products", headers=admin_token).json()
        category_id = next(p["category_id"] for p in products if p.get("category_id"))
        r = requests.delete(f"{BASE}/api/admin/categories/{category_id}", headers=admin_token)
        assert r.status_code == 409

    def test_admin_create_product_missing_fields(self, admin_token):
        # 缺必填字段应 422
        r = requests.post(f"{BASE}/api/admin/products", json={"sku_code": "X"}, headers=admin_token)
        assert r.status_code == 422

    def test_admin_ship_unknown_order(self, admin_token):
        r = requests.post(f"{BASE}/api/admin/orders/ORD_NOT_EXIST/ship", headers=admin_token)
        assert r.status_code == 404

    # ---------- 用户管理 ----------
    def test_admin_list_customers(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/customers", headers=admin_token)
        assert r.status_code == 200
        customers = r.json()
        assert isinstance(customers, list)
        # buyer 用户应存在且含统计字段
        buyer = next((c for c in customers if c["email"] == BUYER_EMAIL), None)
        if buyer:
            for key in ("orders_count", "total_spent", "is_active"):
                assert key in buyer

    def test_admin_customers_search(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/customers", params={"q": BUYER_EMAIL}, headers=admin_token)
        assert r.status_code == 200
        assert any(c["email"] == BUYER_EMAIL for c in r.json())

    def test_admin_toggle_customer(self, admin_token):
        # 先查 buyer id
        customers = requests.get(f"{BASE}/api/admin/customers", params={"q": BUYER_EMAIL}, headers=admin_token).json()
        buyer = next((c for c in customers if c["email"] == BUYER_EMAIL), None)
        if not buyer:
            pytest.skip("buyer 用户不存在")
        cid = buyer["id"]
        # 禁用
        r = requests.put(f"{BASE}/api/admin/customers/{cid}/status",
                         json={"is_active": False}, headers=admin_token)
        assert r.status_code == 200, r.text
        assert r.json()["is_active"] is False
        # 恢复
        r2 = requests.put(f"{BASE}/api/admin/customers/{cid}/status",
                          json={"is_active": True}, headers=admin_token)
        assert r2.status_code == 200
        assert r2.json()["is_active"] is True

    # ---------- 商品删除 ----------
    def test_admin_delete_product(self, admin_token):
        sku_code = f"PYTEST-DEL-{ts}"
        payload = {
            "sku_code": sku_code,
            "name_zh": "待删除商品",
            "base_price": "9.99",
            "skus": [{"sku_code": f"{sku_code}-S", "price": "9.99", "stock": 1}],
        }
        created = requests.post(f"{BASE}/api/admin/products", json=payload, headers=admin_token)
        assert created.status_code == 201, created.text
        pid = created.json()["id"]

        # 删除为「软删除」：商品移入回收站，仍可查到（tab=deleted）
        r = requests.delete(f"{BASE}/api/admin/products/{pid}", headers=admin_token)
        assert r.status_code == 200, r.text
        assert "回收站" in r.json()["message"]
        in_bin = requests.get(
            f"{BASE}/api/admin/products?tab=deleted&product_id={pid}", headers=admin_token
        ).json()
        assert [p["id"] for p in in_bin] == [pid]
        # 默认列表（tab=all）不再包含
        in_all = requests.get(
            f"{BASE}/api/admin/products?product_id={pid}", headers=admin_token
        ).json()
        assert in_all == []

        # 重复删除幂等（仍返回 200，提示已在回收站）
        r2 = requests.delete(f"{BASE}/api/admin/products/{pid}", headers=admin_token)
        assert r2.status_code == 200
        assert "回收站" in r2.json()["message"]

        # 彻底删除后彻底查不到
        r3 = requests.delete(f"{BASE}/api/admin/products/{pid}/purge", headers=admin_token)
        assert r3.status_code == 200, r3.text
        assert requests.get(
            f"{BASE}/api/admin/products?tab=deleted&product_id={pid}", headers=admin_token
        ).json() == []

    # ---------- 管理员账号管理 + 权限 ----------
    def test_admin_manage_admins(self, admin_token):
        uname = f"ptest_{ts}"
        payload = {"username": uname, "password": "pass123456", "full_name": "测试管理员", "role": "operator"}
        r = requests.post(f"{BASE}/api/admin/admins", json=payload, headers=admin_token)
        assert r.status_code == 201, r.text
        aid = r.json()["id"]
        # 列表包含
        admins = requests.get(f"{BASE}/api/admin/admins", headers=admin_token).json()
        assert any(a["id"] == aid for a in admins)
        # 改角色
        r2 = requests.put(f"{BASE}/api/admin/admins/{aid}", json={"role": "viewer"}, headers=admin_token)
        assert r2.status_code == 200
        assert r2.json()["role"] == "viewer"
        # 删除
        r3 = requests.delete(f"{BASE}/api/admin/admins/{aid}", headers=admin_token)
        assert r3.status_code == 200

    def test_admin_role_permission_viewer_readonly(self, admin_token):
        # 创建 viewer 并登录，验证只读
        uname = f"ptest_v_{ts}"
        requests.post(f"{BASE}/api/admin/admins",
                      json={"username": uname, "password": "pass123456", "role": "viewer"}, headers=admin_token)
        login = requests.post(f"{BASE}/api/admin/login", json={"username": uname, "password": "pass123456"})
        assert login.status_code == 200
        vtoken = {"Authorization": f"Bearer {login.json()['access_token']}", **H}
        # 读操作允许
        assert requests.get(f"{BASE}/api/admin/products", headers=vtoken).status_code == 200
        # 写操作禁止
        r = requests.post(f"{BASE}/api/admin/products",
                          json={"sku_code": "NO-PERM", "name_zh": "x", "base_price": 1}, headers=vtoken)
        assert r.status_code == 403
        # 管理员列表禁止
        assert requests.get(f"{BASE}/api/admin/admins", headers=vtoken).status_code == 403
        # 清理
        admins = requests.get(f"{BASE}/api/admin/admins", headers=admin_token).json()
        for a in admins:
            if a["username"] == uname:
                requests.delete(f"{BASE}/api/admin/admins/{a['id']}", headers=admin_token)

    # ---------- 修改 / 重置密码 ----------
    def test_change_own_password_rejects_bad_input(self, admin_token):
        """非法输入必须被拒且**不能改动密码**（沿用固定 admin123 供其余用例登录）"""
        def put(body, token=admin_token):
            return requests.put(f"{BASE}/api/admin/me/password", json=body, headers=token)

        # 当前密码错误
        r = put({"old_password": "wrong-password", "new_password": "brandnew123"})
        assert r.status_code == 400 and "当前密码" in r.json()["detail"]
        # 新旧相同
        r = put({"old_password": ADMIN_PASS, "new_password": ADMIN_PASS})
        assert r.status_code == 400 and "不能与当前密码相同" in r.json()["detail"]
        # 新密码过短 → pydantic 校验
        assert put({"old_password": ADMIN_PASS, "new_password": "123"}).status_code == 422
        # 缺字段
        assert put({"new_password": "brandnew123"}).status_code == 422
        # 未登录
        assert requests.put(f"{BASE}/api/admin/me/password",
                            json={"old_password": ADMIN_PASS,
                                  "new_password": "brandnew123"}).status_code == 401
        # 密码未被改动
        assert requests.post(f"{BASE}/api/admin/login",
                             json={"username": ADMIN_USER,
                                   "password": ADMIN_PASS}).status_code == 200

    def test_superadmin_reset_own_password_is_blocked(self, admin_token):
        """不允许通过「重置他人密码」接口绕过原密码校验改自己"""
        me = requests.get(f"{BASE}/api/admin/admins", headers=admin_token).json()
        my_id = [a["id"] for a in me if a["username"] == ADMIN_USER][0]
        r = requests.put(f"{BASE}/api/admin/admins/{my_id}/password",
                         json={"new_password": "bypass123"}, headers=admin_token)
        assert r.status_code == 400 and "修改密码" in r.json()["detail"]
        # 原密码仍可用
        assert requests.post(f"{BASE}/api/admin/login",
                             json={"username": ADMIN_USER,
                                   "password": ADMIN_PASS}).status_code == 200

    def test_password_change_and_reset_flow(self, admin_token):
        """完整流程：改自己密码 → 登录态切换 → 超管重置 → 权限校验 → 清理"""
        uname = f"pwtest_{ts}"
        created = requests.post(
            f"{BASE}/api/admin/admins",
            json={"username": uname, "password": "init123456",
                  "full_name": "密码测试", "role": "operator"},
            headers=admin_token,
        )
        assert created.status_code == 201, created.text
        aid = created.json()["id"]
        login_url = f"{BASE}/api/admin/login"
        try:
            # 自己的旧密码登录
            me = requests.post(login_url, json={"username": uname, "password": "init123456"})
            assert me.status_code == 200
            mtoken = {"Authorization": f"Bearer {me.json()['access_token']}", **H}

            # operator 修改自己的密码（需原密码）
            r = requests.put(f"{BASE}/api/admin/me/password",
                             json={"old_password": "init123456",
                                   "new_password": "self123456"}, headers=mtoken)
            assert r.status_code == 200, r.text
            # 旧密码失效、新密码可登录
            assert requests.post(login_url, json={"username": uname,
                                                  "password": "init123456"}).status_code == 401
            assert requests.post(login_url, json={"username": uname,
                                                  "password": "self123456"}).status_code == 200

            # operator 无权重置他人密码（即便有合法 token）
            assert requests.put(f"{BASE}/api/admin/admins/{aid}/password",
                                json={"new_password": "hacked123"},
                                headers=mtoken).status_code == 403

            # 超管重置该账号密码（无需对方原密码）
            r = requests.put(f"{BASE}/api/admin/admins/{aid}/password",
                             json={"new_password": "reset12345"}, headers=admin_token)
            assert r.status_code == 200, r.text
            assert "已重置" in r.json()["message"]
            assert requests.post(login_url, json={"username": uname,
                                                  "password": "self123456"}).status_code == 401
            assert requests.post(login_url, json={"username": uname,
                                                  "password": "reset12345"}).status_code == 200

            # 不存在 / 过短密码
            assert requests.put(f"{BASE}/api/admin/admins/999999999/password",
                                json={"new_password": "whatever1"},
                                headers=admin_token).status_code == 404
            assert requests.put(f"{BASE}/api/admin/admins/{aid}/password",
                                json={"new_password": "123"},
                                headers=admin_token).status_code == 422
            # 未登录 401
            assert requests.put(f"{BASE}/api/admin/admins/{aid}/password",
                                json={"new_password": "whatever1"}).status_code == 401
        finally:
            requests.delete(f"{BASE}/api/admin/admins/{aid}", headers=admin_token)

    def test_legacy_admin_update_cannot_change_own_password(self, admin_token):
        """PUT /admins/{id} 的 password 字段只能改他人，改自己需走 /me/password"""
        admins = requests.get(f"{BASE}/api/admin/admins", headers=admin_token).json()
        my_id = [a["id"] for a in admins if a["username"] == ADMIN_USER][0]
        r = requests.put(f"{BASE}/api/admin/admins/{my_id}",
                         json={"password": "bypass123"}, headers=admin_token)
        assert r.status_code == 400 and "修改密码" in r.json()["detail"]
        assert requests.post(f"{BASE}/api/admin/login",
                             json={"username": ADMIN_USER,
                                   "password": ADMIN_PASS}).status_code == 200


# ---------- 后台订单查询（增强）----------
def _pick_sku() -> tuple[int, dict]:
    """取一个有库存的在售商品与首个 SKU"""
    products = requests.get(f"{BASE}/api/products", headers=H).json()
    pid = next(p["id"] for p in products)
    detail = requests.get(f"{BASE}/api/products/{pid}", headers=H).json()
    return pid, detail["skus"][0]


def _checkout(buyer_token, sku: dict, *, quantity: int = 1, receiver_name: str = "订单测试",
              receiver_phone: str = "13800001111",
              receiver_address: str = "上海市浦东新区测试路 1 号",
              remark: str = "pytest 订单管理", pay: bool = False) -> str:
    """购物车加购 → 下单（可选模拟支付），返回订单号"""
    requests.delete(f"{BASE}/api/cart", headers=buyer_token)
    add = requests.post(f"{BASE}/api/cart/items",
                        json={"sku_id": sku["id"], "quantity": quantity},
                        headers=buyer_token)
    assert add.status_code == 201, add.text
    r = requests.post(f"{BASE}/api/orders/checkout", json={
        "receiver_name": receiver_name,
        "receiver_phone": receiver_phone,
        "receiver_address": receiver_address,
        "remark": remark,
        "payment_method": "mock",
    }, headers=buyer_token)
    assert r.status_code == 201, r.text
    order_no = r.json()["order_no"]
    if pay:
        p = requests.post(f"{BASE}/api/orders/{order_no}/pay", headers=buyer_token)
        assert p.status_code == 200, p.text
        c = requests.get(f"{BASE}/api/payments/mock/confirm",
                         params={"txn_no": p.json()["transaction_no"]})
        assert c.status_code == 200, c.text
    return order_no


class TestAdminOrderManagement:
    """订单查询增强接口：状态统计、多条件筛选、发货/备注/取消/收款/完成、批量与导出。"""

    @pytest.fixture()
    def track(self, buyer_token, admin_token):
        """记录本用例创建的订单号，结束后统一清理（避免残留锁库存）"""
        created: list[str] = []
        yield created
        for no in created:
            requests.post(f"{BASE}/api/admin/orders-search/{no}/cancel",
                          json={"reason": "pytest 清理", "refund": True},
                          headers=admin_token)

    def _search(self, admin_token, **params):
        r = requests.get(f"{BASE}/api/admin/orders-search", params=params, headers=admin_token)
        assert r.status_code == 200, r.text
        return r.json()

    def _count(self, admin_token, **params):
        r = requests.get(f"{BASE}/api/admin/orders-search-count", params=params,
                         headers=admin_token)
        assert r.status_code == 200, r.text
        return r.json()["total"]

    # ---------- 鉴权与参数校验 ----------
    def test_endpoints_require_auth(self):
        for path in ("/orders-status-counts", "/orders-search", "/orders-search-count",
                     "/orders-search-export"):
            assert requests.get(f"{BASE}/api/admin{path}").status_code == 401
        for path, body in (
            ("/orders-search/ORD_X/ship", {}),
            ("/orders-search/ORD_X/note", {"note": "x"}),
            ("/orders-search/ORD_X/cancel", {}),
            ("/orders-search/ORD_X/complete", {}),
            ("/orders-search/ORD_X/confirm-payment", {}),
            ("/orders-search/bulk", {"order_nos": ["ORD_X"], "action": "note"}),
        ):
            assert requests.post(f"{BASE}/api/admin{path}", json=body).status_code == 401

    def test_bad_tab_rejected(self, admin_token):
        for params in ({"tab": "not-a-tab"}, {"tab": "unknown"}):
            r = requests.get(f"{BASE}/api/admin/orders-search", params=params,
                             headers=admin_token)
            assert r.status_code == 400, r.text
            r2 = requests.get(f"{BASE}/api/admin/orders-search-count", params=params,
                              headers=admin_token)
            assert r2.status_code == 400

    def test_all_tabs_accepted(self, admin_token):
        """7 个状态标签页都必须可用（曾因 ORDER_TABS 漏了 cancelled 报 400）"""
        for tab in ("all", "pending", "paid", "shipped", "completed",
                    "refunded", "cancelled"):
            r = requests.get(f"{BASE}/api/admin/orders-search",
                             params={"tab": tab, "page_size": 1}, headers=admin_token)
            assert r.status_code == 200, f"tab={tab} 被拒绝：{r.text}"
            c = requests.get(f"{BASE}/api/admin/orders-search-count",
                             params={"tab": tab}, headers=admin_token)
            assert c.status_code == 200, f"tab={tab} 计数被拒绝：{c.text}"

    def test_tabs_cover_all_order_statuses(self, admin_token):
        """各标签页数量之和 = 全部（保证没有状态被漏在标签页之外）"""
        counts = {}
        for tab in ("pending", "paid", "shipped", "completed",
                    "refunded", "cancelled"):
            counts[tab] = self._count(admin_token, tab=tab)
        assert self._count(admin_token, tab="all") == sum(counts.values())

    def test_cancelled_tab_lists_cancelled_order(self, admin_token, buyer_token, track):
        """取消后的订单必须出现在「已取消」标签页（回归：曾报 tab 非法）"""
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku)
        track.append(no)
        requests.post(f"{BASE}/api/admin/orders-search/{no}/cancel",
                      json={"reason": "标签页回归"}, headers=admin_token)

        rows = self._search(admin_token, tab="cancelled", order_no=no)
        assert [o["order_no"] for o in rows] == [no]
        assert rows[0]["status_label"] == "已取消"

    # ---------- 状态统计 ----------
    def test_status_counts_shape(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/orders-status-counts", headers=admin_token)
        assert r.status_code == 200, r.text
        data = r.json()
        for key in ("all", "pending", "paid", "shipped", "completed",
                    "cancelled", "refunded", "pending_ship_alert"):
            assert key in data, key
            assert isinstance(data[key], int) and data[key] >= 0
        parts = sum(data[k] for k in ("pending", "paid", "shipped",
                                      "completed", "cancelled", "refunded"))
        assert data["all"] == parts

    def test_status_counts_reflect_new_order(self, admin_token, buyer_token, track):
        before = requests.get(f"{BASE}/api/admin/orders-status-counts",
                              headers=admin_token).json()
        _, sku = _pick_sku()
        track.append(_checkout(buyer_token, sku))

        after = requests.get(f"{BASE}/api/admin/orders-status-counts",
                             headers=admin_token).json()
        assert after["pending"] == before["pending"] + 1
        assert after["all"] == before["all"] + 1

    # ---------- 筛选 ----------
    def test_search_filters(self, admin_token, buyer_token, track):
        pid, sku = _pick_sku()
        no = _checkout(buyer_token, sku, receiver_name="筛选测试甲",
                       receiver_phone="13711112222", remark="筛选备注ABC")
        track.append(no)
        no2 = _checkout(buyer_token, sku, receiver_name="筛选测试乙",
                        receiver_phone="13733334444")
        track.append(no2)

        # 收件人 / 手机号
        assert {o["order_no"] for o in self._search(admin_token, receiver="筛选测试甲")} == {no}
        assert {o["order_no"] for o in self._search(admin_token, phone="13733334444")} == {no2}
        # 订单号精确（单值走模糊匹配）
        assert {o["order_no"] for o in self._search(admin_token, order_no=no)} == {no}
        # 订单号多值（逗号分隔）
        both = self._search(admin_token, order_no=f"{no},{no2}")
        assert {o["order_no"] for o in both} == {no, no2}
        # 通用关键词命中收件人与备注无关字段（按收件人/手机号匹配）
        assert {o["order_no"] for o in self._search(admin_token, keyword="筛选测试甲")} == {no}
        # 按商品过滤：商品 ID（纯数字）与规格编码两种口径都能命中
        assert no in {o["order_no"] for o in self._search(admin_token, product_id=str(pid))}
        assert no in {o["order_no"] for o in self._search(admin_token,
                                                          product_id=sku["sku_code"])}
        # 不存在的关键词
        assert self._search(admin_token, keyword="ZZZ_NOT_EXIST_QQQ") == []

    def test_search_filter_matches_count(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku, receiver_name="计数一致性")
        track.append(no)

        params = {"receiver": "计数一致性", "page": 1, "page_size": 5}
        rows = self._search(admin_token, **params)
        assert self._count(admin_token, **{"receiver": "计数一致性"}) == len(rows) == 1

    def test_search_tab_filter(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku, pay=True)
        track.append(no)

        rows = self._search(admin_token, tab="paid", order_no=no)
        assert [o["order_no"] for o in rows] == [no]
        assert self._search(admin_token, tab="pending", order_no=no) == []

    def test_search_date_range(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku)
        track.append(no)

        today = time.strftime("%Y-%m-%d")
        assert no in {o["order_no"] for o in self._search(admin_token, date_from=today,
                                                          date_to=today)}
        assert no not in {o["order_no"] for o in self._search(admin_token, date_from="2099-01-01")}
        # 非法日期被忽略而不是报错
        assert self._search(admin_token, date_from="not-a-date", order_no=no)

    def test_search_pagination(self, admin_token):
        rows = self._search(admin_token, page=1, page_size=2)
        assert len(rows) <= 2
        total = self._count(admin_token)
        assert total >= len(rows)
        # 超出范围的页码返回空
        assert self._search(admin_token, page=99999, page_size=20) == []

    def test_search_limit_mode(self, admin_token):
        rows = self._search(admin_token, limit=3)
        assert len(rows) <= 3

    # ---------- 详情 ----------
    def test_order_detail_fields(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku, quantity=2, pay=True)
        track.append(no)

        r = requests.get(f"{BASE}/api/admin/orders-search/{no}", headers=admin_token)
        assert r.status_code == 200, r.text
        o = r.json()
        assert o["order_no"] == no
        assert o["status"] == "paid" and o["status_label"] == "待发货"
        assert o["total_quantity"] == 2 and o["item_count"] == 1
        assert float(o["goods_amount"]) > 0
        assert o["customer_email"] == BUYER_EMAIL
        assert o["items"][0]["sku_id"] == sku["id"]
        assert o["payments"], "已支付订单应有支付记录"
        assert o["payments"][0]["method_label"] == "模拟支付"
        assert o["payments"][0]["status_label"] == "支付成功"
        # 时间节点为本地时间：下单时间不晚于支付时间
        assert o["created_at"] <= o["paid_at"], (o["created_at"], o["paid_at"])

    def test_order_detail_unknown(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/orders-search/ORD_NOT_EXIST", headers=admin_token)
        assert r.status_code == 404

    # ---------- 发货 ----------
    def test_ship_and_repeat(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku, pay=True)
        track.append(no)

        r = requests.post(f"{BASE}/api/admin/orders-search/{no}/ship",
                          json={"carrier": "顺丰速运", "tracking_no": "SF123456"},
                          headers=admin_token)
        assert r.status_code == 200, r.text
        o = r.json()
        assert o["status"] == "shipped" and o["status_label"] == "待收货"
        assert o["carrier"] == "顺丰速运" and o["tracking_no"] == "SF123456"
        assert o["shipped_at"] and o["shipped_at"] >= o["paid_at"]

        again = requests.post(f"{BASE}/api/admin/orders-search/{no}/ship", headers=admin_token)
        assert again.status_code == 400 and "已发货" in again.json()["detail"]

    def test_ship_requires_paid(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku)  # 未支付
        track.append(no)

        r = requests.post(f"{BASE}/api/admin/orders-search/{no}/ship", headers=admin_token)
        assert r.status_code == 400 and "仅已支付订单可发货" in r.json()["detail"]

    def test_ship_unknown_order(self, admin_token):
        r = requests.post(f"{BASE}/api/admin/orders-search/ORD_NOT_EXIST/ship",
                          headers=admin_token)
        assert r.status_code == 404

    def test_tracking_filter_after_ship(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku, pay=True)
        track.append(no)
        requests.post(f"{BASE}/api/admin/orders-search/{no}/ship",
                      json={"tracking_no": "ZT_TEST_0001"}, headers=admin_token)

        assert {o["order_no"] for o in self._search(admin_token, tracking_no="ZT_TEST_0001")} == {no}
        assert {o["order_no"] for o in self._search(admin_token, keyword="ZT_TEST_0001")} == {no}

    # ---------- 备注 ----------
    def test_admin_note(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku)
        track.append(no)

        r = requests.post(f"{BASE}/api/admin/orders-search/{no}/note",
                          json={"note": "客户要求纸质发票"}, headers=admin_token)
        assert r.status_code == 200, r.text
        assert r.json()["admin_note"] == "客户要求纸质发票"

        got = requests.get(f"{BASE}/api/admin/orders-search/{no}", headers=admin_token).json()
        assert got["admin_note"] == "客户要求纸质发票"
        # 清空
        assert requests.post(f"{BASE}/api/admin/orders-search/{no}/note",
                             json={"note": ""}, headers=admin_token).json()["admin_note"] is None

    def test_admin_note_too_long(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku)
        track.append(no)
        r = requests.post(f"{BASE}/api/admin/orders-search/{no}/note",
                          json={"note": "x" * 501}, headers=admin_token)
        assert r.status_code == 422

    # ---------- 取消 / 退款 ----------
    def test_cancel_unpaid_releases_locked_stock(self, admin_token, buyer_token, track):
        pid, sku = _pick_sku()
        before = sku["available_stock"]
        assert before >= 1

        no = _checkout(buyer_token, sku)
        track.append(no)
        after_checkout = requests.get(f"{BASE}/api/products/{pid}", headers=H).json()
        assert after_checkout["skus"][0]["available_stock"] == before - 1

        r = requests.post(f"{BASE}/api/admin/orders-search/{no}/cancel",
                          json={"reason": "买家不要了"}, headers=admin_token)
        assert r.status_code == 200, r.text
        o = r.json()
        assert o["status"] == "cancelled" and o["status_label"] == "已取消"
        assert o["cancel_reason"] == "买家不要了" and o["cancelled_at"]

        after_cancel = requests.get(f"{BASE}/api/products/{pid}", headers=H).json()
        assert after_cancel["skus"][0]["available_stock"] == before

    def test_cancel_paid_requires_refund_and_restores_stock(self, admin_token, buyer_token, track):
        pid, sku = _pick_sku()
        before = sku["available_stock"]
        no = _checkout(buyer_token, sku, pay=True)
        track.append(no)

        paid_stock = requests.get(f"{BASE}/api/products/{pid}", headers=H).json()
        assert paid_stock["skus"][0]["available_stock"] == before - 1

        # 已支付订单不允许直接取消
        bad = requests.post(f"{BASE}/api/admin/orders-search/{no}/cancel",
                            json={"reason": "不想要了"}, headers=admin_token)
        assert bad.status_code == 400 and "退款" in bad.json()["detail"]

        r = requests.post(f"{BASE}/api/admin/orders-search/{no}/cancel",
                          json={"reason": "拍错了", "refund": True}, headers=admin_token)
        assert r.status_code == 200, r.text
        o = r.json()
        assert o["status"] == "refunded" and o["status_label"] == "退款/售后"
        assert o["payments"] and o["payments"][0]["status_label"] == "已退款"

        restored = requests.get(f"{BASE}/api/products/{pid}", headers=H).json()
        assert restored["skus"][0]["available_stock"] == before

    def test_cancel_twice_rejected(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku)
        track.append(no)
        assert requests.post(f"{BASE}/api/admin/orders-search/{no}/cancel",
                             headers=admin_token).status_code == 200
        again = requests.post(f"{BASE}/api/admin/orders-search/{no}/cancel",
                              headers=admin_token)
        assert again.status_code == 400 and "已取消" in again.json()["detail"]

    # ---------- 确认收款 / 完成 ----------
    def test_confirm_payment(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku)  # 未支付
        track.append(no)

        r = requests.post(f"{BASE}/api/admin/orders-search/{no}/confirm-payment",
                          headers=admin_token)
        assert r.status_code == 200, r.text
        o = r.json()
        assert o["status"] == "paid" and o["paid_at"]

        again = requests.post(f"{BASE}/api/admin/orders-search/{no}/confirm-payment",
                              headers=admin_token)
        assert again.status_code == 400 and "待付款" in again.json()["detail"]

    def test_complete_flow_and_then_cancel_rejected(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku, pay=True)
        track.append(no)

        premature = requests.post(f"{BASE}/api/admin/orders-search/{no}/complete",
                                  headers=admin_token)
        assert premature.status_code == 400 and "待收货" in premature.json()["detail"]

        requests.post(f"{BASE}/api/admin/orders-search/{no}/ship", headers=admin_token)
        r = requests.post(f"{BASE}/api/admin/orders-search/{no}/complete",
                          headers=admin_token)
        assert r.status_code == 200, r.text
        o = r.json()
        assert o["status"] == "completed" and o["status_label"] == "已完成"
        assert o["completed_at"] and o["completed_at"] >= o["shipped_at"]

        cancel = requests.post(f"{BASE}/api/admin/orders-search/{no}/cancel",
                               headers=admin_token)
        assert cancel.status_code == 400 and "已完成" in cancel.json()["detail"]

    # ---------- 批量 ----------
    def test_bulk_ship_with_skipped(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        n1 = _checkout(buyer_token, sku, pay=True)
        track.append(n1)
        n2 = _checkout(buyer_token, sku, pay=True)
        track.append(n2)
        n3 = _checkout(buyer_token, sku)  # 未支付，应被跳过
        track.append(n3)

        r = requests.post(f"{BASE}/api/admin/orders-search/bulk", json={
            "order_nos": [n1, n2, n3, "ORD_NOT_EXIST"],
            "action": "ship", "carrier": "圆通速递",
        }, headers=admin_token)
        assert r.status_code == 200, r.text
        msg = r.json()["message"]
        assert "成功 2 条" in msg and "跳过 2 条" in msg
        assert "ORD_NOT_EXIST" in msg

        for no in (n1, n2):
            od = requests.get(f"{BASE}/api/admin/orders-search/{no}",
                              headers=admin_token).json()
            assert od["status"] == "shipped" and od["carrier"] == "圆通速递"

    def test_bulk_note(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        n1 = _checkout(buyer_token, sku)
        track.append(n1)
        n2 = _checkout(buyer_token, sku)
        track.append(n2)

        r = requests.post(f"{BASE}/api/admin/orders-search/bulk", json={
            "order_nos": [n1, n2, n1],  # 重复项应去重
            "action": "note", "note": "批量备注：加急",
        }, headers=admin_token)
        assert r.status_code == 200, r.text
        assert "成功 2 条" in r.json()["message"]
        for no in (n1, n2):
            assert requests.get(f"{BASE}/api/admin/orders-search/{no}",
                                headers=admin_token).json()["admin_note"] == "批量备注：加急"

    def test_bulk_invalid_payload(self, admin_token):
        bad_action = requests.post(f"{BASE}/api/admin/orders-search/bulk", json={
            "order_nos": ["ORD_X"], "action": "delete",
        }, headers=admin_token)
        assert bad_action.status_code in (400, 422)
        empty = requests.post(f"{BASE}/api/admin/orders-search/bulk", json={
            "order_nos": [], "action": "note",
        }, headers=admin_token)
        assert empty.status_code == 422

    # ---------- 导出 ----------
    def test_export_csv(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/orders-search-export", headers=admin_token)
        assert r.status_code == 200, r.text
        assert "text/csv" in r.headers["content-type"]
        assert "attachment" in r.headers.get("content-disposition", "")
        assert r.content[:3] == b"\xef\xbb\xbf", "缺少 UTF-8 BOM"
        text = r.content.decode("utf-8-sig")
        header = text.splitlines()[0]
        for col in ("订单号", "订单状态", "实收金额", "商品件数", "下单时间"):
            assert col in header, col

    def test_export_respects_filters(self, admin_token, buyer_token, track):
        _, sku = _pick_sku()
        no = _checkout(buyer_token, sku, receiver_name="导出专用收件人")
        track.append(no)

        r = requests.get(f"{BASE}/api/admin/orders-search-export",
                         params={"receiver": "导出专用收件人"}, headers=admin_token)
        assert r.status_code == 200, r.text
        lines = r.content.decode("utf-8-sig").strip().splitlines()
        assert len(lines) == 2, lines
        assert no in lines[1]


class TestBanner:
    """首页轮播（home_hero）管理：独立编辑页依赖的 CRUD 与单条接口。"""

    def test_banner_crud_and_get_one(self, admin_token):
        # 创建
        payload = {
            "placement": "home_hero",
            "title_i18n": {"zh": "测试轮播", "en": "Test Banner"},
            "subtitle_i18n": {"zh": "副标题", "en": "Sub"},
            "button_text_i18n": {"zh": "去看看", "en": "Explore"},
            "image_url": "/static/uploads/test.png",
            "video_url": None,
            "link_url": "/products.html",
            "sort_order": 99,
            "is_active": True,
        }
        r = requests.post(f"{BASE}/api/admin/banners", json=payload, headers=admin_token)
        assert r.status_code == 201, r.text
        banner = r.json()
        bid = banner["id"]

        # 单条详情接口（独立编辑页回填用）
        one = requests.get(f"{BASE}/api/admin/banners/{bid}", headers=admin_token)
        assert one.status_code == 200, one.text
        assert one.json()["id"] == bid
        assert one.json()["title_i18n"]["zh"] == "测试轮播"
        assert one.json()["link_url"] == "/products.html"
        assert one.json()["button_text_i18n"]["en"] == "Explore"

        # 列表包含
        lst = requests.get(f"{BASE}/api/admin/banners?placement=home_hero", headers=admin_token).json()
        assert any(b["id"] == bid for b in lst)

        # 未登录访问单条 401
        assert requests.get(f"{BASE}/api/admin/banners/{bid}").status_code == 401

        # 更新
        upd = requests.put(f"{BASE}/api/admin/banners/{bid}", json={
            "title_i18n": {"zh": "测试轮播改", "en": "Updated"},
            "is_active": False,
        }, headers=admin_token)
        assert upd.status_code == 200, upd.text
        assert upd.json()["title_i18n"]["zh"] == "测试轮播改"
        assert upd.json()["is_active"] is False

        # 公开接口不返回停用轮播
        pub = requests.get(f"{BASE}/api/banners?placement=home_hero").json()
        assert not any(b["id"] == bid for b in pub)

        # 单条不存在 404
        assert requests.get(f"{BASE}/api/admin/banners/999999999", headers=admin_token).status_code == 404

        # 删除
        d = requests.delete(f"{BASE}/api/admin/banners/{bid}", headers=admin_token)
        assert d.status_code == 200, d.text
        # 删除后单条 404
        assert requests.get(f"{BASE}/api/admin/banners/{bid}", headers=admin_token).status_code == 404


def _make_test_pptx(path, sku_code: str, name_zh: str = "蜜桃测试文胸") -> None:
    """生成一份最小产品册 PPT（1 页 = 1 商品），用于导入接口测试。

    页面结构刻意模仿真实产品册：主信息框内「颜色/Color」之后跟颜色文字。
    """
    from io import BytesIO

    from PIL import Image
    from pptx import Presentation as PptxPresentation
    from pptx.util import Inches

    prs = PptxPresentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式

    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(5), Inches(4))
    box.text_frame.text = (
        f"货号/Number\n{sku_code}\n\n"
        f"品名：{name_zh}\n"
        "面料：82.5%聚酯纤维17.5%氨纶\n"
        "尺码：S/M/L\n\n"
        "颜色/Color\n"
        "星耀黑          春雨绿"
    )

    buf = BytesIO()
    Image.new("RGB", (600, 800), (200, 120, 140)).save(buf, format="JPEG")
    slide.shapes.add_picture(BytesIO(buf.getvalue()), Inches(6), Inches(0.5),
                             Inches(3), Inches(4))

    prs.save(str(path))


class TestProductImport:
    """产品册 PPT 批量导入接口（/api/admin/import/*）。"""

    def test_import_format_requires_auth(self):
        assert requests.get(f"{BASE}/api/admin/import/format").status_code == 401

    def test_import_format_rules(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/import/format", headers=admin_token)
        assert r.status_code == 200, r.text
        data = r.json()
        assert ".pptx" in data["supported_ext"]
        assert data["one_slide_per_product"] is True
        assert {"sku_code", "name", "fabrics", "sizes", "colors", "images"} <= set(data["fields"])
        codes = {c["code"] for c in data["categories"]}
        assert {"bras", "vests", "tshirts", "pants", "other"} <= codes
        assert {"merge", "replace"} <= set(data["modes"])

    def test_import_requires_auth(self, tmp_path):
        pptx = tmp_path / f"noauth_{ts}.pptx"
        _make_test_pptx(pptx, f"PTIMP{ts % 1000000}")
        with open(pptx, "rb") as f:
            r = requests.post(f"{BASE}/api/admin/import/products",
                              files={"file": (pptx.name, f)}, data={"dry_run": "true"})
        assert r.status_code == 401

    def test_import_rejects_non_pptx(self, admin_token, tmp_path):
        bogus = tmp_path / f"bogus_{ts}.txt"
        bogus.write_text("not a pptx", encoding="utf-8")
        with open(bogus, "rb") as f:
            r = requests.post(f"{BASE}/api/admin/import/products",
                              files={"file": (bogus.name, f)},
                              data={"dry_run": "true"}, headers=admin_token)
        assert r.status_code == 400, r.text

    def test_import_rejects_bad_mode(self, admin_token, tmp_path):
        pptx = tmp_path / f"badmode_{ts}.pptx"
        _make_test_pptx(pptx, f"PTIMP{ts % 1000000}")
        with open(pptx, "rb") as f:
            r = requests.post(f"{BASE}/api/admin/import/products",
                              files={"file": (pptx.name, f)},
                              data={"mode": "upsert"}, headers=admin_token)
        assert r.status_code == 400, r.text

    def test_import_dry_run_parses_without_writing(self, admin_token, tmp_path):
        """dry_run 应解析出商品字段，但不写库、返回 imported=0。"""
        sku = f"PTIMP{ts % 1000000}"
        pptx = tmp_path / f"dryrun_{ts}.pptx"
        _make_test_pptx(pptx, sku)

        before = requests.get(f"{BASE}/api/admin/dashboard", headers=admin_token).json()["products_count"]

        with open(pptx, "rb") as f:
            r = requests.post(
                f"{BASE}/api/admin/import/products",
                files={"file": (pptx.name, f)},
                data={"dry_run": "true", "preview_limit": "1"},
                headers=admin_token,
            )
        assert r.status_code == 200, r.text
        d = r.json()

        assert d["dry_run"] is True
        assert d["total_slides"] == 1
        assert d["parsed_products"] == 1
        assert d["imported"] == 0
        assert d["products_total"] == 0

        item = d["preview"][0]
        assert item["sku_code"] == sku
        assert item["name_zh"] == "蜜桃测试文胸"
        assert item["category_code"] == "bras"
        assert item["sizes"] == ["S", "M", "L"]
        assert "星耀黑" in item["colors"] and "春雨绿" in item["colors"]
        assert item["main_image"].endswith("/main.jpg")

        # 预览不应改变库中商品数
        after = requests.get(f"{BASE}/api/admin/dashboard", headers=admin_token).json()["products_count"]
        assert after == before


class TestProductReviewAndBulk:
    """商品人工审核、批量管理、审计日志与后台接口文档。

    所有写操作都作用于本测试自建的 `PYTEST-` 前缀商品，
    不触碰真实商品数据。
    """

    @staticmethod
    def _make_product(admin_token, suffix: str) -> int:
        """创建一个测试商品，返回其 id"""
        sku = f"PYTEST-{suffix}-{ts}"
        r = requests.post(
            f"{BASE}/api/admin/products",
            json={
                "sku_code": sku,
                "name_zh": f"审核测试商品 {suffix}",
                "name_en": f"Review Test {suffix}",
                "base_price": 199,
                "status": "active",
                "skus": [{"sku_code": f"{sku}-001", "price": 199, "stock": 10}],
            },
            headers=admin_token,
        )
        assert r.status_code == 201, r.text
        return r.json()["id"]

    # ---------- 审核状态 ----------
    def test_review_stats_shape(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/products-review-stats", headers=admin_token)
        assert r.status_code == 200, r.text
        d = r.json()
        for k in ("total", "pending", "approved", "rejected", "imported", "manual"):
            assert k in d, f"缺少字段 {k}"
        assert d["total"] == d["pending"] + d["approved"] + d["rejected"]

    def test_review_stats_requires_auth(self):
        assert requests.get(f"{BASE}/api/admin/products-review-stats").status_code == 401

    def test_manual_product_is_approved_by_default(self, admin_token):
        """后台人工录入的商品应直接为已通过、来源 manual"""
        pid = self._make_product(admin_token, "manual")
        r = requests.get(f"{BASE}/api/admin/products?q=PYTEST-manual", headers=admin_token)
        item = [p for p in r.json() if p["id"] == pid][0]
        assert item["review_status"] == "approved"
        assert item["source"] == "manual"
        assert item["reviewed_by"] == "admin"

    def test_review_flow(self, admin_token):
        """驳回（需理由）→ 重置待审 → 通过"""
        pid = self._make_product(admin_token, "flow")
        url = f"{BASE}/api/admin/products/{pid}/review"

        # 驳回缺理由 → 400
        assert requests.post(url, json={"action": "reject"}, headers=admin_token).status_code == 400

        # 驳回带理由 → 成功，且理由落库
        r = requests.post(url, json={"action": "reject", "note": "图片不清晰"},
                          headers=admin_token)
        assert r.status_code == 200, r.text
        assert r.json()["review_status"] == "rejected"
        assert r.json()["review_note"] == "图片不清晰"

        # 重置为待审核
        r = requests.post(url, json={"action": "pending", "note": ""}, headers=admin_token)
        assert r.json()["review_status"] == "pending"

        # 通过
        r = requests.post(url, json={"action": "approve", "note": "OK"}, headers=admin_token)
        assert r.json()["review_status"] == "approved"

        # 非法 action → 400
        assert requests.post(url, json={"action": "nope"},
                             headers=admin_token).status_code == 400
        # 不存在的商品 → 404
        assert requests.post(f"{BASE}/api/admin/products/99999999/review",
                             json={"action": "approve"},
                             headers=admin_token).status_code == 404

    def test_pending_product_hidden_from_storefront(self, admin_token):
        """未过审的商品不应出现在前台列表与详情"""
        pid = self._make_product(admin_token, "hidden")
        # 先确认前台可见（人工录入默认 approved）
        assert requests.get(f"{BASE}/api/products/{pid}").status_code == 200
        # 驳回后前台不可见
        requests.post(f"{BASE}/api/admin/products/{pid}/review",
                      json={"action": "reject", "note": "测试下架"},
                      headers=admin_token)
        assert requests.get(f"{BASE}/api/products/{pid}").status_code == 404
        ids = [p["id"] for p in requests.get(
            f"{BASE}/api/products?page_size=100").json()]
        assert pid not in ids

    # ---------- 批量操作 ----------
    def test_bulk_status_and_featured(self, admin_token):
        ids = [self._make_product(admin_token, f"st{i}") for i in range(3)]
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "status", "value": "draft"},
                          headers=admin_token)
        assert r.status_code == 200, r.text
        assert r.json()["succeeded"] == 3
        assert r.json()["failed"] == 0

        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "featured", "value": True},
                          headers=admin_token)
        assert r.json()["succeeded"] == 3

        got = requests.get(f"{BASE}/api/admin/products?q=PYTEST-st", headers=admin_token).json()
        for p in got:
            if p["id"] in ids:
                assert p["status"] == "draft"
                assert p["is_featured"] is True

    def test_bulk_status_writes_and_clears_off_shelf_reason(self, admin_token):
        """批量改状态带 note 时必须写入「下架原因」，重新上架时清空

        曾漏传 note：后台弹窗让用户填下架原因，但请求体里没有该字段，
        原因被静默丢弃（off_shelf_reason 一直为空）。
        """
        pid = self._make_product(admin_token, "ofs")

        r = requests.post(f"{BASE}/api/admin/products/bulk", json={
            "ids": [pid], "action": "status", "value": "off_shelf",
            "note": "批量下架原因",
        }, headers=admin_token)
        assert r.status_code == 200, r.text
        assert r.json()["succeeded"] == 1
        got = requests.get(f"{BASE}/api/admin/products/{pid}", headers=admin_token).json()
        assert got["status"] == "off_shelf"
        assert got["off_shelf_reason"] == "批量下架原因"
        assert got["display_status_label"] == "已下架"

        # 重新上架 → 原因清空（与单个下架/上架接口语义一致）
        requests.post(f"{BASE}/api/admin/products/bulk", json={
            "ids": [pid], "action": "status", "value": "active",
        }, headers=admin_token)
        got3 = requests.get(f"{BASE}/api/admin/products/{pid}", headers=admin_token).json()
        assert got3["status"] == "active" and got3["off_shelf_reason"] is None

    def test_bulk_review(self, admin_token):
        ids = [self._make_product(admin_token, f"rv{i}") for i in range(2)]
        # 批量驳回缺理由 → 400
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "review", "value": "reject"},
                          headers=admin_token)
        assert r.status_code == 400
        # 批量驳回带理由
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "review", "value": "reject",
                                "note": "批量驳回原因"},
                          headers=admin_token)
        assert r.json()["succeeded"] == 2
        # 批量通过
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "review", "value": "approve"},
                          headers=admin_token)
        assert r.json()["succeeded"] == 2
        got = requests.get(f"{BASE}/api/admin/products?q=PYTEST-rv", headers=admin_token).json()
        assert all(p["review_status"] == "approved" for p in got if p["id"] in ids)

    def test_bulk_category(self, admin_token):
        pid = self._make_product(admin_token, "cat")
        cats = requests.get(f"{BASE}/api/admin/categories", headers=admin_token).json()
        assert cats, "需要至少一个分类"
        cid = cats[0]["id"]
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": [pid], "action": "category", "category_id": cid},
                          headers=admin_token)
        assert r.json()["succeeded"] == 1
        got = requests.get(f"{BASE}/api/admin/products?q=PYTEST-cat", headers=admin_token).json()
        assert [p for p in got if p["id"] == pid][0]["category_id"] == cid

    def test_bulk_delete(self, admin_token):
        pid = self._make_product(admin_token, "del")
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": [pid], "action": "delete"}, headers=admin_token)
        assert r.json()["succeeded"] == 1
        assert requests.get(f"{BASE}/api/admin/products?q=PYTEST-del",
                            headers=admin_token).json() == []
        # 删除后审计日志仍保留（product_id 置空，靠货号追溯）
        logs = requests.get(
            f"{BASE}/api/admin/product-audit-logs?q=PYTEST-del", headers=admin_token
        ).json()
        assert any(l["action"] in ("create", "delete", "bulk_delete") for l in logs)

    def test_bulk_validation(self, admin_token):
        pid = self._make_product(admin_token, "valid")
        # 非法 action → 400
        assert requests.post(f"{BASE}/api/admin/products/bulk",
                             json={"ids": [pid], "action": "nope"},
                             headers=admin_token).status_code == 400
        # 非法 status 值 → 400
        assert requests.post(f"{BASE}/api/admin/products/bulk",
                             json={"ids": [pid], "action": "status", "value": "bad"},
                             headers=admin_token).status_code == 400
        # 非法 review 值 → 400
        assert requests.post(f"{BASE}/api/admin/products/bulk",
                             json={"ids": [pid], "action": "review", "value": "bad"},
                             headers=admin_token).status_code == 400
        # 空 ids → 422（Pydantic 校验）
        assert requests.post(f"{BASE}/api/admin/products/bulk",
                             json={"ids": [], "action": "delete"},
                             headers=admin_token).status_code == 422
        # 不存在的 id → 计入 skipped 而非失败
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": [99999999], "action": "status", "value": "active"},
                          headers=admin_token)
        assert r.json()["skipped"] == 1

    def test_bulk_requires_auth(self):
        assert requests.post(f"{BASE}/api/admin/products/bulk",
                             json={"ids": [1], "action": "delete"}).status_code == 401

    def test_viewer_cannot_bulk_write(self, admin_token):
        """只读角色不能执行批量写操作"""
        uname = f"ptest_pv_{ts}"
        requests.post(f"{BASE}/api/admin/admins",
                      json={"username": uname, "password": "pass123456", "role": "viewer"},
                      headers=admin_token)
        login = requests.post(f"{BASE}/api/admin/login",
                              json={"username": uname, "password": "pass123456"})
        vtoken = {"Authorization": f"Bearer {login.json()['access_token']}", **H}
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": [1], "action": "status", "value": "draft"},
                          headers=vtoken)
        assert r.status_code == 403
        # 清理测试管理员
        for a in requests.get(f"{BASE}/api/admin/admins", headers=admin_token).json():
            if a["username"] == uname:
                requests.delete(f"{BASE}/api/admin/admins/{a['id']}", headers=admin_token)

    # ---------- 审计日志 ----------
    def test_audit_logs_recorded(self, admin_token):
        pid = self._make_product(admin_token, "log")
        # 改价 → 产生 update 记录
        requests.put(f"{BASE}/api/admin/products/{pid}", json={"base_price": 259},
                     headers=admin_token)
        logs = requests.get(f"{BASE}/api/admin/products/{pid}/audit-logs",
                            headers=admin_token).json()
        actions = [l["action"] for l in logs]
        assert "create" in actions and "update" in actions
        assert all(l["operator"] == "admin" for l in logs)
        upd = [l for l in logs if l["action"] == "update"][0]
        assert "fields" in upd["detail"] and "base_price" in upd["detail"]["fields"]

    def test_audit_logs_filters_and_auth(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/product-audit-logs?limit=10",
                         headers=admin_token)
        assert r.status_code == 200
        assert isinstance(r.json(), list)
        # 按动作过滤
        r2 = requests.get(f"{BASE}/api/admin/product-audit-logs?action=create&limit=5",
                          headers=admin_token)
        assert all(l["action"] == "create" for l in r2.json())
        # 未登录 401
        assert requests.get(f"{BASE}/api/admin/product-audit-logs").status_code == 401

    # ---------- 商品列表筛选 ----------
    def test_product_list_filters(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/products?review_status=approved&limit=5",
                         headers=admin_token)
        assert r.status_code == 200
        assert all(p["review_status"] == "approved" for p in r.json())

        r = requests.get(f"{BASE}/api/admin/products?status=draft&limit=5",
                         headers=admin_token)
        assert all(p["status"] == "draft" for p in r.json())

        r = requests.get(f"{BASE}/api/admin/products?source=manual&limit=5",
                         headers=admin_token)
        assert all(p["source"] == "manual" for p in r.json())

    # ---------- 接口文档 ----------
    def test_api_docs(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/api-docs/products", headers=admin_token)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["auth"]["type"]
        assert d["review_workflow"]["field"] == "review_status"
        groups = {g["group"] for g in d["groups"]}
        assert {"商品查询（状态标签页 / 多字段搜索）", "商品增删改（人工管理）",
                "上下架与状态流转", "回收站（软删除）", "行内快捷编辑",
                "批量管理", "人工审核", "操作审计"} <= groups
        paths = [i["path"] for g in d["groups"] for i in g["items"]]
        assert "/api/admin/products/bulk" in paths
        assert "/api/admin/products/{id}/review" in paths
        assert "/api/admin/products/{id}/publish" in paths
        assert "/api/admin/products/{id}/purge" in paths
        assert "/api/admin/products/{id}/price" in paths
        # 商品状态派生说明
        assert d["product_status"]["values"]["sold_out"]

    def test_api_docs_requires_auth(self):
        assert requests.get(f"{BASE}/api/admin/api-docs/products").status_code == 401


class TestProductLifecycle:
    """商品生命周期（PDD 风格）：状态标签页、上下架校验、回收站、行内改价改库存。

    写操作均作用于测试自建商品（`PYTEST-` 前缀），不污染真实数据。
    """

    @staticmethod
    def _make(admin_token, suffix: str, *, stock: int = 10, image: bool = True) -> int:
        sku = f"PYTEST-LC-{suffix}-{ts}"
        body = {
            "sku_code": sku,
            "name_zh": f"生命周期测试 {suffix}",
            "base_price": 199,
            "status": "active",
            "main_image": "/static/uploads/jyt/s1/main.jpg" if image else None,
            "skus": [{"sku_code": f"{sku}-001", "price": 199, "stock": stock}],
        }
        r = requests.post(f"{BASE}/api/admin/products", json=body, headers=admin_token)
        assert r.status_code == 201, r.text
        return r.json()["id"]

    def _get(self, admin_token, pid):
        """按 ID 取商品（跨「全部」与「回收站」两个标签查找）"""
        for tab in ("all", "deleted"):
            r = requests.get(
                f"{BASE}/api/admin/products?tab={tab}&product_id={pid}",
                headers=admin_token,
            )
            assert r.status_code == 200, r.text
            found = [p for p in r.json() if p["id"] == pid]
            if found:
                return found[0]
        raise AssertionError(f"商品 {pid} 未查询到")

    # ---------- 状态标签页与计数 ----------
    def test_status_counts(self, admin_token):
        r = requests.get(f"{BASE}/api/admin/products-status-counts", headers=admin_token)
        assert r.status_code == 200, r.text
        d = r.json()
        for k in ("all", "on_sale", "off_shelf", "sold_out",
                  "pending", "rejected", "draft", "deleted"):
            assert k in d, f"缺少字段 {k}"
            assert d[k] >= 0

    def test_status_counts_requires_auth(self):
        assert requests.get(f"{BASE}/api/admin/products-status-counts").status_code == 401

    def test_tab_filter_and_pagination(self, admin_token):
        # 分页形状
        r = requests.get(f"{BASE}/api/admin/products?tab=all&page=1&page_size=5",
                         headers=admin_token)
        assert r.status_code == 200, r.text
        assert len(r.json()) <= 5

        # 非法 tab → 400
        assert requests.get(f"{BASE}/api/admin/products?tab=nope",
                            headers=admin_token).status_code == 400

    def test_search_by_id_and_sku_code(self, admin_token):
        pid = self._make(admin_token, "search")
        # 多 ID 查询（逗号分隔）
        r = requests.get(f"{BASE}/api/admin/products?product_id={pid},99999999",
                         headers=admin_token)
        ids = [p["id"] for p in r.json()]
        assert pid in ids and 99999999 not in ids
        # 非法 ID → 400
        assert requests.get(f"{BASE}/api/admin/products?product_id=abc",
                            headers=admin_token).status_code == 400
        # 规格编码查询
        sku = self._get(admin_token, pid)["sku_code"]
        r = requests.get(f"{BASE}/api/admin/products?sku_code={sku}-001",
                         headers=admin_token)
        assert any(p["id"] == pid for p in r.json())

    # ---------- 状态派生 ----------
    def test_display_status_derivation(self, admin_token):
        pid = self._make(admin_token, "status")
        assert self._get(admin_token, pid)["display_status"] == "on_sale"

        # 库存清零 → 已售罄
        requests.put(f"{BASE}/api/admin/products/{pid}/stock",
                     json={"mode": "set", "value": 0}, headers=admin_token)
        assert self._get(admin_token, pid)["display_status"] == "sold_out"

    # ---------- 上下架 ----------
    def test_publish_check_and_publish(self, admin_token):
        pid = self._make(admin_token, "pub")
        chk = requests.get(f"{BASE}/api/admin/products/{pid}/publish-check",
                           headers=admin_token)
        assert chk.status_code == 200
        assert chk.json()["can_publish"] is True

        # 下架
        r = requests.post(f"{BASE}/api/admin/products/{pid}/unpublish",
                          json={"reason": "测试下架"}, headers=admin_token)
        assert r.status_code == 200, r.text
        assert r.json()["display_status"] == "off_shelf"
        assert r.json()["off_shelf_reason"] == "测试下架"

        # 上架
        r = requests.post(f"{BASE}/api/admin/products/{pid}/publish",
                          headers=admin_token)
        assert r.json()["display_status"] == "on_sale"

    def test_publish_blocked_without_image(self, admin_token):
        pid = self._make(admin_token, "noimg", image=False)
        r = requests.post(f"{BASE}/api/admin/products/{pid}/unpublish",
                          json={}, headers=admin_token)
        assert r.status_code == 200
        # 无主图 → 上架被拦
        chk = requests.get(f"{BASE}/api/admin/products/{pid}/publish-check",
                           headers=admin_token).json()
        assert chk["can_publish"] is False
        assert any("主图" in b for b in chk["blockers"])

        r = requests.post(f"{BASE}/api/admin/products/{pid}/publish", headers=admin_token)
        assert r.status_code == 400
        assert "blockers" in r.json()["detail"]

    def test_publish_blocked_when_rejected(self, admin_token):
        pid = self._make(admin_token, "rej")
        requests.post(f"{BASE}/api/admin/products/{pid}/review",
                      json={"action": "reject", "note": "驳回测试"}, headers=admin_token)
        chk = requests.get(f"{BASE}/api/admin/products/{pid}/publish-check",
                           headers=admin_token).json()
        assert chk["can_publish"] is False
        assert any("驳回" in b for b in chk["blockers"])

    # ---------- 回收站 ----------
    def test_soft_delete_restore_purge(self, admin_token):
        pid = self._make(admin_token, "bin")

        # 软删除
        r = requests.delete(f"{BASE}/api/admin/products/{pid}", headers=admin_token)
        assert r.status_code == 200
        item = self._get(admin_token, pid)
        assert item["display_status"] == "deleted"
        assert item["deleted_at"] is not None

        # 回收站商品不出现在 all 标签
        ids = [p["id"] for p in requests.get(
            f"{BASE}/api/admin/products?tab=all&page_size=200", headers=admin_token).json()]
        assert pid not in ids
        # deleted 标签能看到
        ids = [p["id"] for p in requests.get(
            f"{BASE}/api/admin/products?tab=deleted&page_size=200", headers=admin_token).json()]
        assert pid in ids
        # 前台不可见
        assert requests.get(f"{BASE}/api/products/{pid}").status_code == 404

        # 回收站中不能下架
        assert requests.post(f"{BASE}/api/admin/products/{pid}/unpublish",
                             json={}, headers=admin_token).status_code == 400

        # 恢复 → 已下架
        r = requests.post(f"{BASE}/api/admin/products/{pid}/restore", headers=admin_token)
        assert r.status_code == 200
        assert r.json()["display_status"] == "off_shelf"
        # 重复恢复 → 400
        assert requests.post(f"{BASE}/api/admin/products/{pid}/restore",
                             headers=admin_token).status_code == 400

        # 彻底删除（需先在回收站）
        assert requests.delete(f"{BASE}/api/admin/products/{pid}/purge",
                               headers=admin_token).status_code == 400
        requests.delete(f"{BASE}/api/admin/products/{pid}", headers=admin_token)
        r = requests.delete(f"{BASE}/api/admin/products/{pid}/purge", headers=admin_token)
        assert r.status_code == 200
        assert requests.get(f"{BASE}/api/admin/products?product_id={pid}",
                            headers=admin_token).json() == []

    def test_soft_delete_is_idempotent(self, admin_token):
        pid = self._make(admin_token, "idem")
        r1 = requests.delete(f"{BASE}/api/admin/products/{pid}", headers=admin_token)
        r2 = requests.delete(f"{BASE}/api/admin/products/{pid}", headers=admin_token)
        assert r1.status_code == 200 and r2.status_code == 200
        assert "回收站" in r2.json()["message"]

    # ---------- 行内改价 / 改库存 ----------
    def test_quick_price(self, admin_token):
        pid = self._make(admin_token, "price")
        r = requests.put(f"{BASE}/api/admin/products/{pid}/price",
                         json={"base_price": 88.5, "sync_skus": True, "reason": "促销"},
                         headers=admin_token)
        assert r.status_code == 200, r.text
        d = r.json()
        assert float(d["base_price"]) == 88.5
        assert all(float(s["price"]) == 88.5 for s in d["skus"])
        # 价格必须 > 0
        assert requests.put(f"{BASE}/api/admin/products/{pid}/price",
                            json={"base_price": 0}, headers=admin_token).status_code == 422

    def test_quick_stock_total_semantics(self, admin_token):
        """改库存按「商品总库存」语义：set 后总量应精确等于目标值"""
        pid = self._make(admin_token, "stock")
        # 再加一个 SKU，制造多规格场景
        sku = self._get(admin_token, pid)["sku_code"]
        requests.post(f"{BASE}/api/admin/products/{pid}/skus",
                      json={"sku_code": f"{sku}-002", "price": 199, "stock": 3},
                      headers=admin_token)

        r = requests.put(f"{BASE}/api/admin/products/{pid}/stock",
                         json={"mode": "set", "value": 100}, headers=admin_token)
        assert r.status_code == 200, r.text
        assert r.json()["total_stock"] == 100

        r = requests.put(f"{BASE}/api/admin/products/{pid}/stock",
                         json={"mode": "add", "value": 30}, headers=admin_token)
        assert r.json()["total_stock"] == 130

        # 减到负数 → 400
        assert requests.put(f"{BASE}/api/admin/products/{pid}/stock",
                            json={"mode": "add", "value": -999},
                            headers=admin_token).status_code == 400
        # set 负数 → 400
        assert requests.put(f"{BASE}/api/admin/products/{pid}/stock",
                            json={"mode": "set", "value": -1},
                            headers=admin_token).status_code == 400
        # 非法 mode → 400
        assert requests.put(f"{BASE}/api/admin/products/{pid}/stock",
                            json={"mode": "mul", "value": 2},
                            headers=admin_token).status_code == 400

    # ---------- 规格（SKU）级改价 / 改库存（不含拼单价）----------
    def test_sku_prices_batch(self, admin_token):
        """按规格批量改价：只改提交的规格，基础价跟随最低规格价"""
        pid = self._make(admin_token, "skup")
        p = self._get(admin_token, pid)
        sku_id = p["skus"][0]["id"]                      # 初始 price=199
        # 再加一个不同价的 SKU
        requests.post(f"{BASE}/api/admin/products/{pid}/skus",
                      json={"sku_code": "PYTEST-SKUP-002", "price": 120, "stock": 5},
                      headers=admin_token)

        r = requests.put(f"{BASE}/api/admin/products/{pid}/sku-prices",
                         json={"items": [{"sku_id": sku_id, "price": 88.5}],
                               "reason": "规格促销"},
                         headers=admin_token)
        assert r.status_code == 200, r.text
        d = r.json()
        changed = {s["id"]: float(s["price"]) for s in d["skus"]}
        assert changed[sku_id] == 88.5
        # 未提交的另一个规格保持原价 120
        other = [s for s in d["skus"] if s["id"] != sku_id][0]
        assert float(other["price"]) == 120
        # 基础价自动跟随最低启用规格价
        assert float(d["base_price"]) == 88.5

    def test_sku_prices_validation(self, admin_token):
        pid = self._make(admin_token, "skupv")
        sku_id = self._get(admin_token, pid)["skus"][0]["id"]
        # 价格必须 > 0 → 422
        assert requests.put(f"{BASE}/api/admin/products/{pid}/sku-prices",
                            json={"items": [{"sku_id": sku_id, "price": 0}]},
                            headers=admin_token).status_code == 422
        # 不属于该商品的 SKU → 400
        other = self._make(admin_token, "skupv2")
        other_sku = self._get(admin_token, other)["skus"][0]["id"]
        assert requests.put(f"{BASE}/api/admin/products/{pid}/sku-prices",
                            json={"items": [{"sku_id": other_sku, "price": 10}]},
                            headers=admin_token).status_code == 400

    def test_sku_stocks_batch(self, admin_token):
        """按规格批量改库存：add 增减 / set 设为，逐 SKU 写入库存流水"""
        pid = self._make(admin_token, "skus")
        p = self._get(admin_token, pid)
        sku_id = p["skus"][0]["id"]                       # 初始 stock=10

        # mode=add：+5
        r = requests.put(f"{BASE}/api/admin/products/{pid}/sku-stocks",
                         json={"mode": "add",
                               "items": [{"sku_id": sku_id, "value": 5}],
                               "reason": "补货"},
                         headers=admin_token)
        assert r.status_code == 200, r.text
        assert r.json()["total_stock"] == 15

        # mode=set：设为 30（覆盖，而不是累加）
        r = requests.put(f"{BASE}/api/admin/products/{pid}/sku-stocks",
                         json={"mode": "set", "items": [{"sku_id": sku_id, "value": 30}]},
                         headers=admin_token)
        assert r.json()["total_stock"] == 30

        # 减成负数 → 400（当前 30，减 999）
        assert requests.put(f"{BASE}/api/admin/products/{pid}/sku-stocks",
                            json={"mode": "add",
                                  "items": [{"sku_id": sku_id, "value": -999}]},
                            headers=admin_token).status_code == 400
        # 非法 mode → 400
        assert requests.put(f"{BASE}/api/admin/products/{pid}/sku-stocks",
                            json={"mode": "mul", "items": [{"sku_id": sku_id, "value": 1}]},
                            headers=admin_token).status_code == 400

        # 库存流水：每次有变更都写一条，且带规格信息字段
        rows = requests.get(f"{BASE}/api/admin/products/{pid}/stock-movements",
                            headers=admin_token).json()
        assert len(rows) >= 2
        m = rows[0]                                    # 最新一条：set 30
        assert m["sku_id"] == sku_id and m["balance_after"] == 30
        assert m["sku_code"]                            # 规格编码
        assert "attributes" in m and isinstance(m["attributes"], dict)
        assert m["reason"] and m["operator"] == "admin"

    def test_sku_stock_add_or_delta_inconsistent(self, admin_token):
        """add 批量时各规格增减量不一致：value 不记录单一口径，摘要走合计"""
        pid = self._make(admin_token, "skud")
        p = self._get(admin_token, pid)
        skus = p["skus"]
        body = {
            "sku_code": "PYTEST-SKUD-002", "price": 199, "stock": 5,
            "attributes": {"颜色": "白", "尺码": "L"},
        }
        requests.post(f"{BASE}/api/admin/products/{pid}/skus", json=body,
                      headers=admin_token)
        items = [{"sku_id": skus[0]["id"], "value": 3},
                 {"sku_id": self._get(admin_token, pid)["skus"][1]["id"], "value": -2}]
        r = requests.put(f"{BASE}/api/admin/products/{pid}/sku-stocks",
                         json={"mode": "add", "items": items}, headers=admin_token)
        assert r.status_code == 200, r.text

        logs = requests.get(f"{BASE}/api/admin/products/{pid}/audit-logs",
                            headers=admin_token).json()
        sku_stock = [l for l in logs if l["action"] == "sku_stock"]
        assert sku_stock
        # 摘要应能渲染，不抛错、不含裸 JSON 花括号
        assert "{" not in sku_stock[0]["summary"]
        assert "规格" in sku_stock[0]["summary"]

    # ---------- 批量状态流转 ----------
    def test_bulk_publish_and_unpublish(self, admin_token):
        ids = [self._make(admin_token, f"bp{i}") for i in range(2)]
        # 先下架
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "unpublish", "note": "批量下架"},
                          headers=admin_token)
        assert r.json()["succeeded"] == 2
        # 批量上架（含体检）
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "publish"}, headers=admin_token)
        assert r.json()["succeeded"] == 2
        for pid in ids:
            assert self._get(admin_token, pid)["display_status"] == "on_sale"

    def test_bulk_publish_reports_blockers(self, admin_token):
        """体检不通过的商品应计入 failed 并给出原因"""
        ok = self._make(admin_token, "bok")
        bad = self._make(admin_token, "bbad", image=False)
        requests.post(f"{BASE}/api/admin/products/{bad}/unpublish",
                      json={}, headers=admin_token)
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": [ok, bad], "action": "publish"}, headers=admin_token)
        d = r.json()
        assert d["succeeded"] == 1 and d["failed"] == 1
        assert any("主图" in e for e in d["errors"])

    def test_bulk_delete_restore_purge(self, admin_token):
        ids = [self._make(admin_token, f"bb{i}") for i in range(2)]
        # 移入回收站
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "delete"}, headers=admin_token)
        assert r.json()["succeeded"] == 2
        for pid in ids:
            assert self._get(admin_token, pid)["display_status"] == "deleted"
        # 重复删除 → skipped
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "delete"}, headers=admin_token)
        assert r.json()["skipped"] == 2
        # 恢复
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "restore"}, headers=admin_token)
        assert r.json()["succeeded"] == 2
        # 未在回收站的商品不能 purge → skipped
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "purge"}, headers=admin_token)
        assert r.json()["succeeded"] == 0 and r.json()["skipped"] == 2
        # 再删除后 purge
        requests.post(f"{BASE}/api/admin/products/bulk",
                      json={"ids": ids, "action": "delete"}, headers=admin_token)
        r = requests.post(f"{BASE}/api/admin/products/bulk",
                          json={"ids": ids, "action": "purge"}, headers=admin_token)
        assert r.json()["succeeded"] == 2
        assert requests.get(f"{BASE}/api/admin/products?product_id={','.join(map(str, ids))}",
                            headers=admin_token).json() == []

    # ---------- 审计日志覆盖新动作 ----------
    def test_audit_logs_cover_lifecycle(self, admin_token):
        pid = self._make(admin_token, "audit")
        requests.post(f"{BASE}/api/admin/products/{pid}/unpublish",
                      json={"reason": "审计测试"}, headers=admin_token)
        requests.post(f"{BASE}/api/admin/products/{pid}/publish", headers=admin_token)
        requests.put(f"{BASE}/api/admin/products/{pid}/price",
                     json={"base_price": 77}, headers=admin_token)
        requests.put(f"{BASE}/api/admin/products/{pid}/stock",
                     json={"mode": "set", "value": 12}, headers=admin_token)
        requests.delete(f"{BASE}/api/admin/products/{pid}", headers=admin_token)

        logs = requests.get(f"{BASE}/api/admin/products/{pid}/audit-logs",
                            headers=admin_token).json()
        actions = {l["action"] for l in logs}
        assert {"create", "unpublish", "publish",
                "quick_price", "quick_stock", "delete"} <= actions

    # ---------- 操作日志展示字段（中文动作 / 摘要 / 配色）----------
    def test_audit_log_display_fields(self, admin_token):
        """日志必须带中文动作标签、徽标色调与人话摘要，不再裸露 JSON"""
        pid = self._make(admin_token, "disp")
        requests.post(f"{BASE}/api/admin/products/{pid}/unpublish",
                      json={"reason": "展示测试"}, headers=admin_token)
        requests.post(f"{BASE}/api/admin/products/{pid}/publish", headers=admin_token)
        requests.put(f"{BASE}/api/admin/products/{pid}/price",
                     json={"base_price": 88.5, "sync_skus": True}, headers=admin_token)
        requests.put(f"{BASE}/api/admin/products/{pid}/stock",
                     json={"mode": "set", "value": 8}, headers=admin_token)
        requests.delete(f"{BASE}/api/admin/products/{pid}", headers=admin_token)

        logs = requests.get(f"{BASE}/api/admin/products/{pid}/audit-logs",
                            headers=admin_token).json()
        by_action = {l["action"]: l for l in logs}
        for log in logs:
            assert log["action_label"]          # 中文动作标签非空
            assert log["tone"] in ("green", "blue", "orange", "red", "gray")
            assert log["summary"] and log["summary"] != "-"

        assert by_action["create"]["action_label"] == "新增商品"
        assert by_action["publish"]["tone"] == "green"
        assert by_action["delete"]["tone"] == "red"

        # 摘要里不应出现裸 JSON 的花括号
        unpub = by_action["unpublish"]
        assert "{" not in unpub["summary"] and "}" not in unpub["summary"]
        assert "→" in unpub["summary"]                    # 已上架？→ 已下架
        assert "原因：展示测试" in unpub["summary"]
        assert "¥199.00 → ¥88.50" in by_action["quick_price"]["summary"]
        assert "总库存 设为 8" in by_action["quick_stock"]["summary"]

    def test_audit_logs_pagination(self, admin_token):
        """分页：offset/limit 生效，且不重复不遗漏"""
        pid = self._make(admin_token, "page")
        for price in range(60, 66):
            requests.put(f"{BASE}/api/admin/products/{pid}/price",
                         json={"base_price": price}, headers=admin_token)

        p1 = requests.get(f"{BASE}/api/admin/products/{pid}/audit-logs?limit=3",
                          headers=admin_token).json()
        p2 = requests.get(f"{BASE}/api/admin/products/{pid}/audit-logs?limit=3&offset=3",
                          headers=admin_token).json()
        assert len(p1) == 3 and len(p2) == 3
        assert not ({l["id"] for l in p1} & {l["id"] for l in p2})   # 无重叠
        assert p1[0]["id"] > p2[0]["id"]                            # 倒序

        # 按动作过滤
        only = requests.get(
            f"{BASE}/api/admin/products/{pid}/audit-logs?action=quick_price",
            headers=admin_token).json()
        assert only and all(l["action"] == "quick_price" for l in only)

    def test_audit_meta_endpoint(self, admin_token):
        """筛选元信息：总数、各动作数量（不受 action 过滤影响）、操作人列表"""
        pid = self._make(admin_token, "meta")
        requests.put(f"{BASE}/api/admin/products/{pid}/price",
                     json={"base_price": 66}, headers=admin_token)

        meta = requests.get(f"{BASE}/api/admin/product-audit-logs/meta?product_id={pid}",
                            headers=admin_token).json()
        assert meta["total"] >= 2
        counts = {a["action"]: a["count"] for a in meta["actions"]}
        assert counts.get("quick_price") == 1
        assert counts.get("create") == 1
        assert all(a["label"] for a in meta["actions"])     # 每个动作都有中文标签
        assert any(a["action"] == "quick_price" and a["label"] == "修改价格"
                   for a in meta["actions"])
        assert "admin" in meta["operators"]

        # 选中某动作后 total 收窄，但动作计数仍保留全部（供下拉展示数量）
        meta2 = requests.get(
            f"{BASE}/api/admin/product-audit-logs/meta?product_id={pid}&action=quick_price",
            headers=admin_token).json()
        assert meta2["total"] == 1
        counts2 = {a["action"]: a["count"] for a in meta2["actions"]}
        assert counts2.get("create") == 1

        # 未登录 401
        assert requests.get(
            f"{BASE}/api/admin/product-audit-logs/meta").status_code == 401

    # ---------- 按 ID 精确加载（编辑页修复回归）----------
    def test_load_single_product_by_id(self, admin_token):
        """`product_id` 精确查询：老商品也应能查到（不被第一页分页截断）

        回归：编辑页原先用不带参数的 /api/admin/products（只返回第一页 20 条），
        导致 ID 靠前的商品无法编辑并报「商品不存在」。
        """
        pid = self._make(admin_token, "byid")
        r = requests.get(f"{BASE}/api/admin/products?product_id={pid}&limit=1",
                         headers=admin_token)
        assert r.status_code == 200, r.text
        items = r.json()
        assert len(items) == 1 and items[0]["id"] == pid
        assert items[0]["skus"]                      # 带 SKU 明细，编辑页可用

        # 对比：不带参数只返回第一页（默认 20 条）
        page1 = requests.get(f"{BASE}/api/admin/products", headers=admin_token).json()
        assert len(page1) <= 20

    def test_admin_product_detail_any_status(self, admin_token):
        """后台单商品详情：不限状态（草稿/下架/回收站都能取），未登录 401"""
        pid = self._make(admin_token, "detail")
        # 下架（前台接口此时会 404）
        requests.post(f"{BASE}/api/admin/products/{pid}/unpublish",
                      json={}, headers=admin_token)

        r = requests.get(f"{BASE}/api/admin/products/{pid}", headers=admin_token)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["id"] == pid
        assert d["display_status"] == "off_shelf"
        assert d["display_name"]                       # 按语言解析出的展示名
        assert d["skus"]

        # 前台接口此时应拿不到
        assert requests.get(f"{BASE}/api/products/{pid}").status_code == 404

        # 未登录后台 → 401；不存在 → 404
        assert requests.get(f"{BASE}/api/admin/products/{pid}").status_code == 401
        assert requests.get(f"{BASE}/api/admin/products/99999999",
                            headers=admin_token).status_code == 404

    def test_products_count_matches_filters(self, admin_token):
        """筛选后的总数接口必须与列表口径一致（分页显示「共 N 条」用）

        回归：后台列表原先用「标签页数量」当总数，筛选后仍显示全量条数、分页也算错。
        """
        pid = self._make(admin_token, "cnt")
        code = self._get(admin_token, pid)["sku_code"]

        # 无筛选：等于该标签页数量
        base = requests.get(f"{BASE}/api/admin/products-count?tab=all",
                            headers=admin_token).json()["total"]
        counts = requests.get(f"{BASE}/api/admin/products-status-counts",
                              headers=admin_token).json()
        assert base == counts["all"]

        # 按编码精确筛选：计数必须与列表返回条数一致
        r = requests.get(f"{BASE}/api/admin/products-count"
                         f"?tab=all&q={code}", headers=admin_token)
        assert r.status_code == 200, r.text
        assert r.json()["total"] == 1
        listed = requests.get(f"{BASE}/api/admin/products?tab=all&q={code}&limit=200",
                              headers=admin_token).json()
        assert len(listed) == 1

        # 命中 0 条
        assert requests.get(f"{BASE}/api/admin/products-count?tab=all&q=ZZZ_NO_HIT_ZZZ",
                            headers=admin_token).json()["total"] == 0

        # 非法 tab → 400
        assert requests.get(f"{BASE}/api/admin/products-count?tab=bad",
                            headers=admin_token).status_code == 400
        # 未登录 → 401
        assert requests.get(f"{BASE}/api/admin/products-count").status_code == 401

    def test_products_detail_images_use_d_prefix(self, admin_token):
        """图文详情图 URL 必须带 /d/ 标记，轮播图不带 —— 前台据此区分两类图"""
        pid = self._make(admin_token, "dmark")
        main = "/static/uploads/jyt/s1/main.jpg"
        detail = "/d/static/uploads/jyt/s1/detail_1.png"
        requests.put(f"{BASE}/api/admin/products/{pid}",
                     json={"images": [main, detail]}, headers=admin_token)

        p = self._get(admin_token, pid)
        gallery = [u for u in p["images"] if "/d/" not in u]
        details = [u for u in p["images"] if "/d/" in u]
        assert gallery == [main]
        assert details == [detail]

        # /d/static/... 与 /static/... 指向同一份静态资源，均可访问
        assert requests.get(f"{BASE}/d/static/uploads/jyt/s1/main.jpg").status_code == 200

    def test_upload_detail_returns_d_prefix(self, admin_token):
        """kind=detail 上传返回带 /d/ 前缀的 URL；kind 非法返回 400"""
        import io

        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (20, 20), (180, 90, 90)).save(buf, format="PNG")
        buf.seek(0)

        r = requests.post(f"{BASE}/api/upload?kind=detail",
                          files={"file": ("d.png", buf, "image/png")},
                          headers=admin_token)
        assert r.status_code == 201, r.text
        assert r.json()["url"].startswith("/d/static/uploads/")
        assert r.json()["kind"] == "detail"

        # 轮播图上传不带 /d/
        buf.seek(0)
        r2 = requests.post(f"{BASE}/api/upload?kind=image",
                           files={"file": ("g.png", buf, "image/png")},
                           headers=admin_token)
        assert r2.json()["url"].startswith("/static/uploads/")
        assert "/d/" not in r2.json()["url"]

        # 非法 kind → 400
        buf.seek(0)
        assert requests.post(f"{BASE}/api/upload?kind=bad",
                             files={"file": ("x.png", buf, "image/png")},
                             headers=admin_token).status_code == 400

    def test_remove_spec_deactivates_skus(self, admin_token):
        """编辑页删除规格后，对应 SKU 应被停用（前台 / 购物车不再可选）

        回归：原先删规格只重建前端矩阵，被移除的 SKU 在库中仍启用。
        """
        pid = self._make(admin_token, "deact")
        base = self._get(admin_token, pid)["sku_code"]
        # 补一个 SKU，形成 2 个规格
        r = requests.post(f"{BASE}/api/admin/products/{pid}/skus",
                          json={"sku_code": f"{base}-B", "price": 199, "stock": 7,
                                "attributes": {"颜色": "白", "尺码": "M"}},
                          headers=admin_token)
        new_sku_id = r.json()["id"]
        before = self._get(admin_token, pid)
        assert before["total_stock"] == 17            # 10 + 7

        # 模拟编辑页「删除规格」：把该 SKU 停用
        requests.put(f"{BASE}/api/admin/skus/{new_sku_id}",
                     json={"sku_code": f"{base}-B", "price": 199, "stock": 7,
                           "attributes": {"颜色": "白", "尺码": "M"},
                           "is_active": False},
                     headers=admin_token)

        after = self._get(admin_token, pid)
        target = [s for s in after["skus"] if s["id"] == new_sku_id][0]
        assert target["is_active"] is False
        assert after["total_stock"] == 10             # 只计启用规格

    def test_add_sku_reuses_same_code_in_product(self, admin_token):
        """同商品内重复 sku_code：复用原 SKU（不报 500）——支持「删规格后重新加回」"""
        pid = self._make(admin_token, "reuse")
        base = self._get(admin_token, pid)["sku_code"]
        code = f"{base}-R"

        r1 = requests.post(f"{BASE}/api/admin/products/{pid}/skus",
                           json={"sku_code": code, "price": 199, "stock": 5,
                                 "attributes": {"颜色": "白"}}, headers=admin_token)
        assert r1.status_code == 201, r1.text
        first_id = r1.json()["id"]

        # 同编码再次提交 → 复用同一条记录并更新（而非 500）
        r2 = requests.post(f"{BASE}/api/admin/products/{pid}/skus",
                           json={"sku_code": code, "price": 88, "stock": 9,
                                 "attributes": {"颜色": "白"}, "is_active": True},
                           headers=admin_token)
        assert r2.status_code == 201, r2.text
        assert r2.json()["id"] == first_id
        assert r2.json().get("reused") is True

        skus = self._get(admin_token, pid)["skus"]
        assert len([s for s in skus if s["sku_code"] == code]) == 1   # 未产生重复记录
        assert float([s for s in skus if s["id"] == first_id][0]["price"]) == 88

    def test_add_sku_code_taken_by_other_product(self, admin_token):
        """跨商品占用 sku_code：返回 400 并给出可读原因（而非 500）"""
        pid1 = self._make(admin_token, "code1")
        pid2 = self._make(admin_token, "code2")
        code = self._get(admin_token, pid1)["skus"][0]["sku_code"]

        r = requests.post(f"{BASE}/api/admin/products/{pid2}/skus",
                          json={"sku_code": code, "price": 10, "stock": 1},
                          headers=admin_token)
        assert r.status_code == 400
        assert "占用" in r.json()["detail"]

        # 更新 SKU 时改用已占编码同样 400
        sku_id = self._get(admin_token, pid2)["skus"][0]["id"]
        r2 = requests.put(f"{BASE}/api/admin/skus/{sku_id}",
                          json={"sku_code": code, "price": 10, "stock": 1,
                                "attributes": {}, "is_active": True},
                          headers=admin_token)
        assert r2.status_code == 400
        assert "占用" in r2.json()["detail"]