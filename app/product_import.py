"""产品册 PPT 导入核心逻辑（CLI 与 HTTP 接口共用）

职责：
1. 解析 pptx：逐页提取 货号/品名/面料/尺码/颜色
2. 提取图片：主图 / 详情图（可选色卡图）保存到 static/uploads/jyt/s{页码}/
3. 写入数据库：Product + SKU（颜色 × 尺码）

被两处调用：
- `app/import_jyt.py`（命令行脚本）
- `app/routers/imports.py`（POST /api/admin/import/products）
"""
from __future__ import annotations

import logging
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Sequence

from pptx import Presentation
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Category, Product, SKU

logger = logging.getLogger(__name__)

# 图片输出根目录（站内 static/uploads 下）
PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUT_BASE = PROJECT_ROOT / "static" / "uploads" / "jyt"
# 图片 URL 前缀（与 main.py 中 /static 挂载一致）
URL_PREFIX = "/static/uploads/jyt"

# ---------------------------------------------------------------- 分类映射
# 按品名关键词归类到服装分类
CATEGORY_RULES = [
    # (分类code, 分类zh, 分类en, 关键词列表)
    ("bras",       "文胸", "Bras",          ["文胸"]),
    ("vests",      "背心", "Vests",         ["背心", "吊带"]),
    ("tshirts",    "短袖", "T-Shirts",      ["短袖"]),
    ("longsleeves", "长袖", "Long Sleeves", ["长袖", "开衫", "罩衫"]),
    ("jackets",    "外套", "Jackets",       ["外套", "连帽衫", "拉链连帽"]),
    ("shorts",     "短裤", "Shorts",        ["短裤", "三分裤", "五分裤", "七分裤"]),
    ("pants",      "长裤", "Pants",         ["长裤", "喇叭裤", "裙裤", "阔腿裤", "瑜伽裤", "机车裤", "烟管裤", "直筒裤"]),
    ("skirts",     "裙子", "Skirts",        ["短裙", "百褶裙"]),
]

# 面料成分基线价格（元）—— 按品类给合理零售价
BASE_PRICE_BY_CAT = {
    "bras": 129, "vests": 99, "tshirts": 119, "longsleeves": 139,
    "jackets": 199, "shorts": 119, "pants": 169, "skirts": 119,
}
DEFAULT_PRICE = 99

# 图片分类阈值
# 1) 版式占位图：PPT 每页都铺了一张大图（占页面 71.7% 或 100%），内容是
#    白底 + 橙色占位线 + 色卡，不含服装，必须剔除，否则会被误当主图。
#    但同样是大图的还有「拼版商品图」（一张图内排了正面/侧面/背面三张模特照），
#    它必须保留。两者面积相同，只能靠内容区分：
#      版式占位图 白底率 ≈ 0.95 ~ 1.00
#      拼版商品图 白底率 ≈ 0.48 ~ 0.73  ← 实测断层在 0.73 与 0.95 之间
#    故规则为「面积大 且 近乎全白」才算占位图。
BACKGROUND_AREA_RATIO = 0.5      # 面积超过页面 50% 才可能是版式大图
BLANK_WHITE_RATIO = 0.85         # 白底率超过该值视为空白版式（无实质内容）
WHITE_PIXEL_THRESHOLD = 235      # 判定「白」的通道阈值
# 2) 色卡/logo 小图：面积低于页面 1.2%
SWATCH_AREA_RATIO = 0.012


def _threshold_white(v: int) -> int:
    """单通道像素值 → 0/255 二值（> WHITE_PIXEL_THRESHOLD 视为白）"""
    return 255 if v > WHITE_PIXEL_THRESHOLD else 0


def is_blank_layout(
    blob: bytes, width: int, height: int, page_area: int, sample: int = 120
) -> bool:
    """判断图片是否为「大面积且近乎全白」的版式占位图（不含服装）。

    仅当面积大（> BACKGROUND_AREA_RATIO）**且**白底率高（> BLANK_WHITE_RATIO）
    时才判定为占位图，避免误删「拼版商品图」。
    任何异常都保守返回 False（保留该图）。
    """
    if page_area <= 0 or (width * height) <= page_area * BACKGROUND_AREA_RATIO:
        return False
    try:
        from io import BytesIO

        from PIL import Image, ImageChops

        im = Image.open(BytesIO(blob)).convert("RGB")
        im.thumbnail((sample, sample))
        # 三通道都判为白才算「白」：各自阈值化成 0/255 掩码后相乘
        # （ImageChops.multiply 按 im1*im2/255 归一，故只有三通道同时为 255 才得 255）
        r, g, b = (c.point(_threshold_white) for c in im.split())
        mask = ImageChops.multiply(ImageChops.multiply(r, g), b)
        white = mask.histogram()[255]
        total = im.width * im.height
        return bool(total) and (white / total) > BLANK_WHITE_RATIO
    except Exception:
        logger.debug("白底率检测失败，保留该图", exc_info=True)
        return False


def pick_category(name: str) -> str:
    """按品名关键词返回分类 code"""
    for code, _, _, kws in CATEGORY_RULES:
        for kw in kws:
            if kw in name:
                return code
    return "other"


# ---------------------------------------------------------------- 字段解析
def slide_texts(slide) -> list[str]:
    """收集一页中所有非空文本框内容"""
    return [
        sh.text_frame.text.strip()
        for sh in slide.shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    ]


def parse_num(texts: list[str], joined: str) -> str | None:
    """提取货号，如 JYJD017 / MTWCX02 / MT5FK"""
    # 1) 独立成行的裸货号（字母开头、含数字，长度 4-12；排除尺码/成分行）
    for t in texts:
        for line in t.splitlines():
            t2 = re.sub(r"[^A-Za-z0-9]", "", line)
            if re.fullmatch(r"[A-Za-z]{2,}[A-Za-z0-9]*\d[A-Za-z0-9]*", t2) \
                    and 4 <= len(t2) <= 12:
                return t2
    # 2) 带"货号/Number"标签（有明确标签，放宽长度上限，避免把长货号截断）
    m = re.search(r"货号/Number\s*[|]?\s*([A-Za-z0-9_-]{2,32})", joined)
    if m:
        return m.group(1)
    # 3) 文本中任意位置
    m = re.search(r"\b([A-Za-z]{2,6}\d{2,})", joined)
    return m.group(1) if m else None


def parse_name(texts: list[str]) -> str:
    """从 "品名：xxx" 行提取中文品名"""
    for t in texts:
        m = re.search(r"品名[：:]\s*([^\n]+)", t)
        if m:
            n = m.group(1).strip().strip("|").strip()
            if n and n != "(无品名)":
                return n
    return "瑜伽服饰"


def parse_fabrics(texts: list[str], joined: str) -> dict[str, str]:
    """解析面料成分：{"面料": "80%锦纶20%氨纶", "网纱": "94%聚酯纤维6%氨纶"}"""
    zh_map = {
        "聚酯纤维": "Polyester", "锦纶": "Nylon", "氨纶": "Elastane",
        "棉": "Cotton", "羊毛": "Wool", "腈纶": "Acrylic",
        "涤纶": "Polyester", "粘纤": "Viscose",
    }
    out: dict[str, str] = {}
    for t in texts:
        # 中文面料行（优先）
        for m in re.finditer(
            r"([\u4e00-\u9fa5]{1,4})[：:]\s*((?:\d+\.?\d*%[\u4e00-\u9fa5A-Za-z/]{1,12})+)",
            t,
        ):
            key, val = m.group(1), m.group(2).strip()
            if key not in ("货号", "品名", "尺码", "颜色", "颜色 ", "Number", "Composition"):
                out[key] = val
    # 无中文面料 → 翻译 Composition 行
    if not out:
        m = re.search(r"Composition:\s*([\d.%A-Za-z, ]+)\b", joined)
        if m:
            comp = m.group(1).strip()
            for zh, en in zh_map.items():
                comp = comp.replace(en, zh)
            out["面料"] = comp
    return out


def parse_sizes(joined: str) -> list[str]:
    """解析尺码行，如 "尺码：S/M/L/XL/XXL" → ["S","M","L","XL","XXL"]"""
    m = re.search(r"尺码[：:]\s*([A-Z]{1,4}(?:/[A-Z]{1,4})+|均码)", joined)
    if not m:
        return ["S", "M", "L", "XL", "XXL"]
    raw = m.group(1)
    if raw == "均码":
        return ["均码"]
    return raw.split("/")


# 英文颜色兜底（页面无中文颜色时使用）
_EN_COLOR_MAP = {
    "Red": "红色", "Blue": "蓝色", "Pink": "粉色", "Purple": "紫色",
    "Grey": "灰色", "Gray": "灰色", "White": "白色", "Black": "黑色",
    "Yellow": "黄色", "Orange": "橙色", "Green": "绿色", "Violet": "紫色",
    "Indigo": "靛蓝色", "Turmeric": "姜黄色", "Apricot": "杏色",
    "Matcha": "抹茶色", "Magenta": "品红色", "Coffee": "咖啡色",
    "Navy": "藏蓝色",
}

# 「颜色/Color」标记（允许字符间夹杂空格/换行，PPT 中常被拆成多行多个 run）
_COLOR_MARKER_RE = re.compile(r"颜\s*色\s*/\s*C\s*o\s*l\s*o\s*r", re.IGNORECASE)
# 非颜色噪音词
_COLOR_NOISE = {
    "颜色", "出货价", "参考", "货号", "品名", "面料", "尺码",
    "正品保障", "极速发货", "全场包邮", "无理由", "退换货",
}
MAX_COLORS = 6


def parse_colors(texts: list[str]) -> list[str]:
    """提取颜色中文名（保序去重，最多 6 个）；无中文时用英文颜色兜底。

    颜色在 PPT 中有两种存放方式：
    1. 独立色卡文本框（整框都是颜色名）
    2. 主信息框内「颜色/Color」标记之后（本函数只取标记之后的部分，
       避免把品名/面料等误当颜色）
    """
    out: list[str] = []
    seen: set[str] = set()

    def add_from(text: str) -> None:
        for zh in re.findall(r"[\u4e00-\u9fa5]{2,6}", text):
            if zh in _COLOR_NOISE or zh in seen:
                continue
            seen.add(zh)
            out.append(zh)

    for t in texts:
        if not t.strip():
            continue
        marker = _COLOR_MARKER_RE.search(t)
        if marker:                       # 有标记 → 只取标记之后的文本
            add_from(t[marker.end():])
            continue
        if "货号" in t or "品名" in t:    # 无标记的主信息框 → 跳过
            continue
        add_from(t)                      # 普通色卡文本框 → 整框提取

    # 英文颜色兜底
    if not out:
        for t in texts:
            if "货号" in t or "品名" in t:
                continue
            for line in t.splitlines():
                for en, zh in _EN_COLOR_MAP.items():
                    if en in line and zh not in seen:
                        seen.add(zh)
                        out.append(zh)
    return out[:MAX_COLORS]


_NAME_MAP = {
    "锦纶": "Nylon", "蜜桃": "Peach", "条纹": "Striped", "宽松": "Loose",
    "文胸": "Sports Bra", "背心": "Tank Top", "短袖": "T-Shirt",
    "长袖": "Long Sleeve", "外套": "Jacket", "短裤": "Shorts",
    "长裤": "Pants", "短裙": "Skirt", "裙裤": "Skirt Pants",
    "网纱": "Mesh", "镂空": "Cutout", "拼色": "Color Block",
    "撞色": "Contrast", "撞条": "Banded", "织带": "Webbing",
    "基础": "Basic", "运动": "Sport",
    "喇叭裤": "Flare Pants", "工字背": "Racerback", "挂脖": "Halter",
    "背交叉": "Crossback", "高腰": "High Waist", "小高领": "Turtleneck",
    "五分裤": "Capri", "七分裤": "7/8 Pants", "阔腿裤": "Wide Leg",
    "连帽": "Hooded", "加绒": "Fleece", "吊带": "Strap",
    "半拉链": "Half Zip", "字母": "Letter", "速干": "Quick Dry",
    "防晒": "UV", "百褶": "Pleated", "连体": "One Piece",
    "抽绳": "Drawstring", "口袋": "Pocket",
    "美背": "Back", "开叉": "Slit", "假两件": "Two-Piece",
    "升级版": "Pro", "拼接": "Patchwork", "卫衣": "Hoodie",
    "喇叭": "Flare", "线感": "Seamed", "菱格": "Diamond",
    "翻腰": "Fold Waist", "抽皱": "Ruffled", "扭结": "Twist",
    "金链": "Gold Chain", "人影线": "Stitch", "束口": "Cuffed",
    "直筒": "Straight", "烟管": "Tapered", "机车": "Moto",
    "甜心": "Sweet", "冰丝": "Ice Silk",
    "弧线": "Curved", "弧背": "Curved Back", "肩带": "Strap",
    "排扣": "Button", "渐变": "Gradient", "吊染": "Dye Ombre",
    "单面": "Single-Faced", "立领": "Stand Collar",
}


def to_en_name(zh: str) -> str:
    """中文品名 → 英文品名（简单映射；未收录词保留原文）"""
    en = zh
    for k, v in _NAME_MAP.items():
        en = en.replace(k, v)
    return en


# ---------------------------------------------------------------- 图片提取
def extract_photos(
    slide,
    slide_no: int,
    page_area: int,
    *,
    save_images: bool = True,
    save_swatches: bool = False,
    out_base: Path | None = None,
    url_prefix: str = URL_PREFIX,
) -> tuple[str | None, list[str], list[str]]:
    """提取一页中的图片并保存（自动剔除版式占位图）。

    分类规则：
    - 「面积 > BACKGROUND_AREA_RATIO 且白底率 > BLANK_WHITE_RATIO」→ 版式占位图，
      **直接丢弃**（白底 + 占位线 + 色卡，不含服装）
    - 剩余图片按面积降序：最大者 → 主图；`<= SWATCH_AREA_RATIO` → 色卡/logo 小图
      （`save_swatches=False` 时不落盘）；其余 → 详情图
    - 若只剩小图（最大者也不超过 SWATCH_AREA_RATIO），说明该页没有服装实拍图，
      返回无主图而不是拿 logo 充数

    返回 (主图URL, 详情图URLs, 小图URLs)
    """
    pics: list[tuple[bytes, str, int]] = []
    for shape in slide.shapes:
        if shape.shape_type != 13:
            continue
        try:
            blob = shape.image.blob
            ext = shape.image.ext
        except Exception:
            continue
        width, height = shape.width, shape.height
        if is_blank_layout(blob, width, height, page_area):
            continue
        pics.append((blob, ext, width * height))
    if not pics:
        return None, [], []
    pics.sort(key=lambda it: it[2], reverse=True)  # 面积降序：第一张为主图

    swatch_area = int(page_area * SWATCH_AREA_RATIO)
    # 剩下的最大一张仍只是 logo/色卡尺寸 → 该页没有服装实拍图，不能拿 logo 当主图
    if pics[0][2] <= swatch_area:
        return None, [], []

    out_dir = (out_base or OUT_BASE) / f"s{slide_no}"

    def save(blob: bytes, ext: str, name: str, *, detail_image: bool = False) -> str:
        if save_images:
            out_dir.mkdir(parents=True, exist_ok=True)
            (out_dir / f"{name}.{ext}").write_bytes(blob)
        # 图文详情图 URL 带 /d/ 前缀标记（主图不带），静态文件由 /d/static 挂载提供，
        # 前台（products.html）与后台编辑页据此把「轮播图 / 图文详情图」区分开
        prefix = f"/d{url_prefix}" if detail_image else url_prefix
        return f"{prefix}/s{slide_no}/{name}.{ext}"

    main_url = save(*pics[0][:2], "main")
    detail_urls: list[str] = []
    swatch_urls: list[str] = []
    di = si = 0
    for blob, ext, area in pics[1:]:
        if area <= swatch_area:
            # 小图多为 logo/色卡，默认不落盘（仅计数），避免产生无用文件
            si += 1
            if save_swatches:
                swatch_urls.append(save(blob, ext, f"swatch_{si}"))
        else:
            di += 1
            detail_urls.append(save(blob, ext, f"detail_{di}", detail_image=True))
    return main_url, detail_urls, swatch_urls


# ---------------------------------------------------------------- 数据结构
@dataclass
class ParsedItem:
    """单个商品页解析结果"""

    slide: int
    sku_code: str
    name_zh: str
    name_en: str
    category_code: str
    base_price: Decimal
    description_zh: str
    description_en: str
    sizes: list[str]
    colors: list[str]
    fabrics: dict[str, str]
    main_image: str | None
    images: list[str]
    warnings: list[str] = field(default_factory=list)


@dataclass
class ParseResult:
    """整个 pptx 的解析结果"""

    source: str
    total_slides: int
    items: list[ParsedItem]
    skipped_slides: list[dict] = field(default_factory=list)
    images: dict[str, int] = field(default_factory=dict)


@dataclass
class ImportSummary:
    """写入数据库后的统计结果"""

    source: str
    total_slides: int
    imported: int
    skipped: int
    failed: int
    products_total: int
    skus_total: int
    categories_total: int
    images: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- 解析 pptx
def build_item(slide_no: int, slide, page_area: int, *, save_images: bool,
               save_swatches: bool, out_base: Path | None, url_prefix: str,
               texts: list[str] | None = None,
               sku_code: str | None = None) -> ParsedItem:
    """把一页 PPT 解析成 ParsedItem

    `texts` / `sku_code` 可由调用方预先算好传入，避免重复解析
    （parse_pptx 需要先拿到货号做去重，再去提取图片）。
    """
    if texts is None:
        texts = slide_texts(slide)
    joined = "\n".join(texts)

    name_zh = parse_name(texts)
    warnings: list[str] = []
    if not sku_code:
        sku_code = parse_num(texts, joined)
    if not sku_code:
        sku_code = f"JYT-{slide_no:03d}"
        warnings.append("未识别到货号，已按页码生成临时货号")

    category_code = pick_category(name_zh)
    base_price = Decimal(str(BASE_PRICE_BY_CAT.get(category_code, DEFAULT_PRICE)))

    # 加绒/外套 → 上调价格
    if "加绒" in name_zh or "加绒" in joined:
        base_price = Decimal(str(BASE_PRICE_BY_CAT.get(category_code, DEFAULT_PRICE) + 40))
    if "外套" in name_zh and "加绒" in joined:
        base_price = Decimal(str(199 + 60))

    fabrics = parse_fabrics(texts, joined)
    sizes = parse_sizes(joined) or ["均码"]
    colors = parse_colors(texts)
    if not colors:
        colors = ["默认色"]
        warnings.append("PPT 中未找到颜色文字（仅色卡图），已使用「默认色」")

    main_image, detail_urls, swatch_urls = extract_photos(
        slide, slide_no, page_area,
        save_images=save_images, save_swatches=save_swatches,
        out_base=out_base, url_prefix=url_prefix,
    )
    if not main_image:
        warnings.append("该页无服装实拍图（仅有版式背景图或完全无图）")
    images = ([main_image] if main_image else []) + detail_urls

    fabric_desc = "；".join(f"{k}：{v}" for k, v in fabrics.items()) or "优质面料"
    en_short = to_en_name(name_zh)
    desc_zh = (
        f"{name_zh}，采用{fabric_desc}，"
        f"提供{'/'.join(sizes)}尺码，"
        f"颜色：{'、'.join(colors)}。"
        f"YOYOLE 瑜伽服系列，柔韧贴合，吸湿速干，适合瑜伽、健身与日常穿搭。"
    )
    desc_en = (
        f"{en_short} in {fabric_desc}, sizes {'/'.join(sizes)}. "
        f"Colors: {', '.join(to_en_name(c) for c in colors)}. "
        f"YOYOLE activewear — soft, stretchy, quick-dry for yoga, "
        f"fitness and daily wear."
    )

    return ParsedItem(
        slide=slide_no,
        sku_code=sku_code,
        name_zh=name_zh,
        name_en=en_short,
        category_code=category_code,
        base_price=base_price,
        description_zh=desc_zh,
        description_en=desc_en,
        sizes=sizes,
        colors=colors,
        fabrics=fabrics,
        main_image=main_image,
        images=images,
        warnings=warnings,
    )


def parse_pptx(
    pptx_path: str | Path,
    *,
    save_images: bool = True,
    save_swatches: bool = False,
    out_base: Path | None = None,
    url_prefix: str = URL_PREFIX,
    clean_out_dir: bool = False,
) -> ParseResult:
    """解析产品册 pptx（不写数据库）。

    - 同一货号出现多页时只保留第一页，其余记入 skipped_slides
    - `save_images=False` 时只解析不落盘（用于预览）
    - `clean_out_dir=True` 时先清空图片输出目录（避免残留旧图）
    """
    pptx_path = Path(pptx_path)
    if not pptx_path.is_file():
        raise FileNotFoundError(f"未找到产品册文件: {pptx_path}")

    out_base = out_base or OUT_BASE
    if save_images and clean_out_dir and out_base.exists():
        import shutil

        shutil.rmtree(out_base, ignore_errors=True)

    prs = Presentation(str(pptx_path))
    slides = list(prs.slides)
    width = prs.slide_width or 0
    height = prs.slide_height or 0
    page_area = int(width) * int(height)

    items: list[ParsedItem] = []
    skipped: list[dict] = []
    seen: dict[str, int] = {}
    img_stats: Counter = Counter()

    for idx, slide in enumerate(slides, 1):
        try:
            # 先取文本与货号做去重判断，避免重复页的图片被白白写盘
            texts = slide_texts(slide)
            sku_code = parse_num(texts, "\n".join(texts)) or f"JYT-{idx:03d}"

            first = seen.get(sku_code)
            if first is not None:
                skipped.append({
                    "slide": idx,
                    "reason": f"货号 {sku_code} 与第 {first} 页重复",
                    "sku_code": sku_code,
                })
                continue

            item = build_item(
                idx, slide, page_area,
                save_images=save_images, save_swatches=save_swatches,
                out_base=out_base, url_prefix=url_prefix,
                texts=texts, sku_code=sku_code,
            )
        except Exception as exc:  # 单页失败不影响整体
            logger.exception("slide %s 解析失败", idx)
            skipped.append({"slide": idx, "reason": f"解析异常: {exc}"})
            continue

        seen[sku_code] = idx
        items.append(item)

        img_stats["main"] += bool(item.main_image)
        img_stats["detail"] += len(item.images) - (1 if item.main_image else 0)

    return ParseResult(
        source=pptx_path.name,
        total_slides=len(slides),
        items=items,
        skipped_slides=skipped,
        images=dict(img_stats),
    )


# ---------------------------------------------------------------- 写入数据库
async def ensure_categories(db: AsyncSession, *, replace: bool) -> dict[str, Category]:
    """保证服装分类存在，返回 {code: Category}；replace=True 时先重建分类表"""
    if replace:
        cats: dict[str, Category] = {}
        for code, zh, en, _ in CATEGORY_RULES:
            cat = Category(code=code, name_i18n={"zh": zh, "en": en},
                           sort_order=len(cats) + 1)
            db.add(cat)
            cats[code] = cat
        other = Category(code="other", name_i18n={"zh": "其他", "en": "Other"},
                         sort_order=99)
        db.add(other)
        cats["other"] = other
        await db.flush()
        return cats

    existing = (await db.execute(select(Category))).scalars().all()
    cats = {c.code: c for c in existing}
    for code, zh, en, _ in CATEGORY_RULES:
        if code not in cats:
            cat = Category(code=code, name_i18n={"zh": zh, "en": en},
                           sort_order=len(cats) + 1)
            db.add(cat)
            cats[code] = cat
    if "other" not in cats:
        other = Category(code="other", name_i18n={"zh": "其他", "en": "Other"},
                         sort_order=99)
        db.add(other)
        cats["other"] = other
    await db.flush()
    return cats


async def import_items(
    db: AsyncSession,
    items: Sequence[ParsedItem],
    *,
    source: str = "",
    mode: str = "replace",
    featured_every: int = 20,
    sku_stock_base: int = 80,
    review_status: str = "pending",
    reviewed_by: str = "import",
) -> ImportSummary:
    """把解析结果写入数据库。

    mode:
    - "replace"：先清空 SKU/Product/Category 再导入（幂等，适合首次或重置）
    - "merge"  ：保留现有数据，仅导入不存在的货号（适合增量补充）

    review_status：导入商品的审核状态（`pending` 待审 / `approved` 直接上架）。
    默认 `pending`，需在后台人工审核后才会在前台展示。
    """
    if mode not in ("replace", "merge"):
        raise ValueError("mode 必须是 replace 或 merge")
    if review_status not in ("pending", "approved", "rejected"):
        raise ValueError("review_status 必须是 pending / approved / rejected")

    if mode == "replace":
        for model in (SKU, Product, Category):
            await db.execute(delete(model))
        await db.flush()

    cats = await ensure_categories(db, replace=mode == "replace")

    existing_nums: set[str] = set()
    if mode == "merge":
        existing_nums = {
            row for row in (await db.execute(select(Product.sku_code))).scalars().all()
        }

    imported = 0
    skipped = 0
    failed = 0
    warnings: list[str] = []
    img_stats: Counter = Counter()

    for item in items:
        if item.sku_code in existing_nums:
            skipped += 1
            warnings.append(f"{item.sku_code} 已存在，已跳过")
            continue
        try:
            cat = cats.get(item.category_code) or cats["other"]
            product = Product(
                category_id=cat.id,
                sku_code=item.sku_code,
                name_i18n={"zh": item.name_zh, "en": item.name_en},
                description_i18n={"zh": item.description_zh, "en": item.description_en},
                main_image=item.main_image,
                images=item.images,
                base_price=item.base_price,
                brand="YOYOLE",
                weight_kg=Decimal("0.300"),
                status="active",
                is_featured=(imported % featured_every == 0) if featured_every else False,
                # 批量导入的商品默认待审核，由后台人工审核后才在前台展示
                source="import",
                review_status=review_status,
                reviewed_by=(
                    reviewed_by
                    if review_status in ("approved", "rejected") else None
                ),
                reviewed_at=(
                    datetime.now(timezone.utc).replace(tzinfo=None)
                    if review_status in ("approved", "rejected") else None
                ),
            )
            db.add(product)
            await db.flush()

            sku_counter = 0
            for color in item.colors:
                for size in item.sizes:
                    sku_counter += 1
                    db.add(
                        SKU(
                            product_id=product.id,
                            sku_code=f"{item.sku_code}-{sku_counter:03d}",
                            attributes={"颜色": color, "尺码": size},
                            price=item.base_price,
                            cost_price=item.base_price * Decimal("0.45"),
                            stock=sku_stock_base + (sku_counter * 13) % 60,
                            locked_stock=0,
                            low_stock_threshold=10,
                            is_active=True,
                        )
                    )
            imported += 1
            existing_nums.add(item.sku_code)
            img_stats["main"] += bool(item.main_image)
            img_stats["detail"] += len(item.images) - (1 if item.main_image else 0)
        except Exception as exc:
            logger.exception("商品 %s 写入失败", item.sku_code)
            failed += 1
            warnings.append(f"{item.sku_code} 写入失败：{exc}")
            await db.rollback()
            cats = await ensure_categories(db, replace=False)

    await db.commit()

    products_total = (await db.execute(select(func.count(Product.id)))).scalar() or 0
    skus_total = (await db.execute(select(func.count(SKU.id)))).scalar() or 0
    categories_total = (await db.execute(select(func.count(Category.id)))).scalar() or 0

    return ImportSummary(
        source=source,
        total_slides=0,
        imported=imported,
        skipped=skipped,
        failed=failed,
        products_total=int(products_total),
        skus_total=int(skus_total),
        categories_total=int(categories_total),
        images=dict(img_stats),
        warnings=warnings,
    )


async def import_pptx(
    db: AsyncSession,
    pptx_path: str | Path,
    *,
    mode: str = "replace",
    save_images: bool = True,
    save_swatches: bool = False,
    clean_out_dir: bool = False,
    featured_every: int = 20,
    review_status: str = "pending",
    reviewed_by: str = "import",
) -> ImportSummary:
    """一步到位：解析 pptx 并写入数据库"""
    parsed = parse_pptx(
        pptx_path, save_images=save_images, save_swatches=save_swatches,
        clean_out_dir=clean_out_dir,
    )
    summary = await import_items(
        db, parsed.items, source=parsed.source, mode=mode,
        featured_every=featured_every, review_status=review_status,
        reviewed_by=reviewed_by,
    )
    summary.total_slides = parsed.total_slides
    summary.skipped += len(parsed.skipped_slides)
    summary.warnings.extend(
        f"slide {s['slide']}: {s['reason']}" for s in parsed.skipped_slides
    )
    for it in parsed.items:
        summary.warnings.extend(f"slide {it.slide}: {w}" for w in it.warnings)
    return summary
